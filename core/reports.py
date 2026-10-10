import io
import csv
from datetime import date, datetime
from decimal import Decimal
from django.db.models import Sum
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Spacer, Paragraph, Table, TableStyle
from .models import Income, Expense, Budget, SavingsGoal, Report, Notification
from .notifications import create_notification


def _normalize_report_month(target_month=None):
    if target_month is None:
        return timezone.localdate().replace(day=1)
    if isinstance(target_month, datetime):
        return target_month.date().replace(day=1)
    if isinstance(target_month, date):
        return target_month.replace(day=1)
    if not isinstance(target_month, str):
        raise ValueError('Month must use YYYY-MM or YYYY-MM-DD format.')

    value = target_month.strip()
    try:
        if len(value) == 7:
            month_date = datetime.strptime(f'{value}-01', '%Y-%m-%d').date()
        elif len(value) == 10:
            month_date = datetime.strptime(value, '%Y-%m-%d').date()
        else:
            raise ValueError
    except ValueError as error:
        raise ValueError('Month must use YYYY-MM or YYYY-MM-DD format.') from error
    return month_date.replace(day=1)


def generate_monthly_report_data(user, target_month=None):
    """
    Generates a full snapshot dictionary of a user's financial data for a given month.
    """
    month_date = _normalize_report_month(target_month)

    month_str = month_date.strftime('%Y-%m')
    month_label = month_date.strftime('%B %Y')

    # Incomes
    incomes = Income.objects.filter(
        user=user,
        date__year=month_date.year,
        date__month=month_date.month
    ).order_by('date', 'id')

    income_transactions = []
    total_income = Decimal('0.00')
    source_totals = {}

    for inc in incomes:
        total_income += inc.amount
        source_totals[inc.source] = source_totals.get(inc.source, Decimal('0.00')) + inc.amount
        income_transactions.append({
            'id': inc.id,
            'date': str(inc.date),
            'source': inc.source,
            'description': inc.description,
            'amount': f'{inc.amount:.2f}',
        })

    # Expenses
    expenses = Expense.objects.filter(
        user=user,
        date__year=month_date.year,
        date__month=month_date.month
    ).order_by('date', 'id')

    expense_transactions = []
    total_expenses = Decimal('0.00')
    category_totals = {}

    for exp in expenses:
        total_expenses += exp.amount
        category_totals[exp.category] = category_totals.get(exp.category, Decimal('0.00')) + exp.amount
        expense_transactions.append({
            'id': exp.id,
            'date': str(exp.date),
            'category': exp.category,
            'description': exp.description,
            'amount': f'{exp.amount:.2f}',
        })

    # Budgets
    budgets = Budget.objects.filter(
        user=user,
        month__year=month_date.year,
        month__month=month_date.month
    )
    budget_map = {b.category: b for b in budgets}
    total_budgeted = Decimal('0.00')
    budget_breakdown = []

    for cat_code, cat_name in Expense.CATEGORY_CHOICES:
        b_obj = budget_map.get(cat_code)
        allocated = b_obj.amount if b_obj else Decimal('0.00')
        spent = category_totals.get(cat_code, Decimal('0.00'))
        remaining = allocated - spent
        util = round(float((spent / allocated) * 100), 2) if allocated > Decimal('0.00') else 0.0

        total_budgeted += allocated
        budget_breakdown.append({
            'category': cat_code,
            'allocated_budget': f'{allocated:.2f}',
            'actual_spent': f'{spent:.2f}',
            'remaining': f'{remaining:.2f}',
            'utilization_percentage': util,
            'is_over_budget': (spent > allocated) if allocated > Decimal('0.00') else (spent > Decimal('0.00')),
        })

    # Savings Goals
    goals = SavingsGoal.objects.filter(user=user)
    goals_data = []
    for g in goals:
        goals_data.append({
            'id': g.id,
            'name': g.name,
            'target_amount': f'{g.target_amount:.2f}',
            'current_amount': f'{g.current_amount:.2f}',
            'progress_percentage': g.progress_percentage,
            'is_completed': g.is_completed,
            'target_date': str(g.target_date) if g.target_date else None,
        })

    net_savings = total_income - total_expenses
    savings_rate = round(float((net_savings / total_income) * 100), 2) if total_income > Decimal('0.00') else 0.0
    overall_budget_util = round(float((total_expenses / total_budgeted) * 100), 2) if total_budgeted > Decimal('0.00') else 0.0

    return {
        'period': month_str,
        'period_label': month_label,
        'user': {
            'username': user.username,
            'email': user.email,
        },
        'summary': {
            'total_income': f'{total_income:.2f}',
            'total_expenses': f'{total_expenses:.2f}',
            'net_savings': f'{net_savings:.2f}',
            'savings_rate': savings_rate,
            'total_budgeted': f'{total_budgeted:.2f}',
            'overall_budget_utilization': overall_budget_util,
            'income_transactions_count': len(income_transactions),
            'expense_transactions_count': len(expense_transactions),
        },
        'budget_breakdown': budget_breakdown,
        'source_breakdown': [
            {'source': src, 'total': f'{tot:.2f}'}
            for src, tot in source_totals.items()
        ],
        'expense_transactions': expense_transactions,
        'income_transactions': income_transactions,
        'savings_goals': goals_data,
    }


def create_monthly_financial_report(user, target_month=None):
    """
    Generates or refreshes a user's monthly report and its system notification.
    """
    report_data = generate_monthly_report_data(user, target_month)
    month_date = datetime.strptime(f"{report_data['period']}-01", '%Y-%m-%d').date()

    title = f"Monthly Financial Report - {report_data['period_label']}"

    report, created = Report.objects.update_or_create(
        user=user,
        report_type='monthly_financial_summary',
        month=month_date,
        defaults={
            'title': title,
            'data': report_data,
            'generated_at': timezone.now(),
        },
    )

    create_notification(
        user=user,
        notification_type='system',
        title=f"Report Ready: {report_data['period_label']}",
        message=f"Your monthly financial summary report for {report_data['period_label']} has been generated.",
        related_context=f"report:{report.id}",
        prevent_duplicate=True
    )

    return report, created


def generate_report_csv_content(report_data):
    """
    Serializes report data into a well-structured CSV string.
    """
    output = io.StringIO()
    writer = csv.writer(output)

    # Header section
    writer.writerow(['BUDGETBUDDY MONTHLY FINANCIAL REPORT'])
    writer.writerow(['Period', report_data.get('period_label', '')])
    writer.writerow(['User', report_data.get('user', {}).get('username', '')])
    writer.writerow(['Generated Date', datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
    writer.writerow([])

    # Summary Section
    summary = report_data.get('summary', {})
    writer.writerow(['=== FINANCIAL SUMMARY ==='])
    writer.writerow(['Metric', 'Value'])
    writer.writerow(['Total Income', f"₹{summary.get('total_income', '0.00')}"])
    writer.writerow(['Total Expenses', f"₹{summary.get('total_expenses', '0.00')}"])
    writer.writerow(['Net Savings', f"₹{summary.get('net_savings', '0.00')}"])
    writer.writerow(['Savings Rate (%)', f"{summary.get('savings_rate', 0)}%"])
    writer.writerow(['Total Budgeted', f"₹{summary.get('total_budgeted', '0.00')}"])
    writer.writerow(['Budget Utilization (%)', f"{summary.get('overall_budget_utilization', 0)}%"])
    writer.writerow([])

    # Budget Breakdown
    writer.writerow(['=== CATEGORY BUDGET ALLOCATION & SPENDING ==='])
    writer.writerow(['Category', 'Allocated Budget (₹)', 'Actual Spent (₹)', 'Remaining (₹)', 'Utilization (%)', 'Over Budget'])
    for b in report_data.get('budget_breakdown', []):
        writer.writerow([
            b.get('category'),
            b.get('allocated_budget'),
            b.get('actual_spent'),
            b.get('remaining'),
            f"{b.get('utilization_percentage')}%",
            'YES' if b.get('is_over_budget') else 'NO',
        ])
    writer.writerow([])

    # Expense Transactions
    writer.writerow(['=== DETAILED EXPENSE TRANSACTIONS ==='])
    writer.writerow(['ID', 'Date', 'Category', 'Description', 'Amount (₹)'])
    for exp in report_data.get('expense_transactions', []):
        writer.writerow([
            exp.get('id'),
            exp.get('date'),
            exp.get('category'),
            exp.get('description') or 'N/A',
            exp.get('amount'),
        ])
    writer.writerow([])

    # Income Transactions
    writer.writerow(['=== DETAILED INCOME TRANSACTIONS ==='])
    writer.writerow(['ID', 'Date', 'Source', 'Description', 'Amount (₹)'])
    for inc in report_data.get('income_transactions', []):
        writer.writerow([
            inc.get('id'),
            inc.get('date'),
            inc.get('source'),
            inc.get('description') or 'N/A',
            inc.get('amount'),
        ])
    writer.writerow([])

    # Savings Goals
    writer.writerow(['=== SAVINGS GOALS SNAPSHOT ==='])
    writer.writerow(['Goal Name', 'Target Amount (₹)', 'Currently Saved (₹)', 'Progress (%)', 'Status', 'Target Date'])
    for g in report_data.get('savings_goals', []):
        writer.writerow([
            g.get('name'),
            g.get('target_amount'),
            g.get('current_amount'),
            f"{g.get('progress_percentage')}%",
            'Completed' if g.get('is_completed') else 'In Progress',
            g.get('target_date') or 'N/A',
        ])

    return output.getvalue()


def generate_report_excel_content(report_data):
    """Serializes report data into a real XLSX workbook."""
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = 'Summary'

    summary_sheet.append(['BUDGETBUDDY MONTHLY FINANCIAL REPORT'])
    summary_sheet.append(['Period', report_data.get('period_label', '')])
    summary_sheet.append(['User', report_data.get('user', {}).get('username', '')])
    summary_sheet.append([])
    summary_sheet.append(['Metric', 'Value'])
    summary = report_data.get('summary', {})
    for label, key in (
        ('Total Income', 'total_income'),
        ('Total Expenses', 'total_expenses'),
        ('Net Savings', 'net_savings'),
        ('Savings Rate (%)', 'savings_rate'),
        ('Total Budgeted', 'total_budgeted'),
        ('Budget Utilization (%)', 'overall_budget_utilization'),
    ):
        summary_sheet.append([label, summary.get(key, 0)])

    def add_table_sheet(title, headers, rows):
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='4F46E5')
        for column in sheet.columns:
            width = min(max(len(str(cell.value or '')) for cell in column) + 2, 36)
            sheet.column_dimensions[column[0].column_letter].width = width
        return sheet

    add_table_sheet(
        'Budgets',
        ['Category', 'Allocated Budget', 'Actual Spent', 'Remaining', 'Utilization (%)', 'Over Budget'],
        [[
            item.get('category'), item.get('allocated_budget'), item.get('actual_spent'),
            item.get('remaining'), item.get('utilization_percentage'), item.get('is_over_budget')
        ] for item in report_data.get('budget_breakdown', [])]
    )
    add_table_sheet(
        'Income',
        ['ID', 'Date', 'Source', 'Description', 'Amount'],
        [[item.get(key) for key in ('id', 'date', 'source', 'description', 'amount')]
         for item in report_data.get('income_transactions', [])]
    )
    add_table_sheet(
        'Expenses',
        ['ID', 'Date', 'Category', 'Description', 'Amount'],
        [[item.get(key) for key in ('id', 'date', 'category', 'description', 'amount')]
         for item in report_data.get('expense_transactions', [])]
    )
    add_table_sheet(
        'Savings Goals',
        ['Name', 'Target Amount', 'Current Amount', 'Progress (%)', 'Completed', 'Target Date'],
        [[item.get(key) for key in ('name', 'target_amount', 'current_amount', 'progress_percentage', 'is_completed', 'target_date')]
         for item in report_data.get('savings_goals', [])]
    )

    summary_sheet.column_dimensions['A'].width = 28
    summary_sheet.column_dimensions['B'].width = 24
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def generate_report_pdf_content(report_data):
    """Serializes report data into a readable PDF document."""
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=0.5 * inch,
        leftMargin=0.5 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )
    styles = getSampleStyleSheet()
    summary = report_data.get('summary', {})
    story = [
        Paragraph('BudgetBuddy Monthly Financial Report', styles['Title']),
        Paragraph(
            f"Period: {report_data.get('period_label', '')} | User: {report_data.get('user', {}).get('username', '')}",
            styles['Normal']
        ),
        Spacer(1, 0.2 * inch),
    ]

    summary_rows = [
        ['Metric', 'Value'],
        ['Total Income', str(summary.get('total_income', '0.00'))],
        ['Total Expenses', str(summary.get('total_expenses', '0.00'))],
        ['Net Savings', str(summary.get('net_savings', '0.00'))],
        ['Total Budgeted', str(summary.get('total_budgeted', '0.00'))],
        ['Budget Utilization (%)', str(summary.get('overall_budget_utilization', 0))],
    ]
    table = Table(summary_rows, colWidths=[2.5 * inch, 2 * inch])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1D5DB')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F4F6')]),
        ('PADDING', (0, 0), (-1, -1), 6),
    ]))
    story.extend([Paragraph('Financial Summary', styles['Heading2']), table, Spacer(1, 0.2 * inch)])

    budget_rows = [['Category', 'Budget', 'Spent', 'Remaining']]
    for item in report_data.get('budget_breakdown', []):
        budget_rows.append([
            item.get('category', ''), item.get('allocated_budget', '0.00'),
            item.get('actual_spent', '0.00'), item.get('remaining', '0.00')
        ])
    budget_table = Table(budget_rows, repeatRows=1)
    budget_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1D5DB')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F4F6')]),
        ('PADDING', (0, 0), (-1, -1), 5),
    ]))
    story.extend([Paragraph('Budget Breakdown', styles['Heading2']), budget_table])
    document.build(story)
    return output.getvalue()
