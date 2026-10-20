from decimal import Decimal
from datetime import date
from io import BytesIO
from django.contrib.auth.models import User
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase
from rest_framework_simplejwt.tokens import RefreshToken

from .models import Expense, Income, Budget, Profile, SavingsGoal, Notification, Report


class HealthCheckTests(APITestCase):
    def test_health_check_is_public_and_checks_database(self):
        response = self.client.get('/api/health/')

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data, {'status': 'ok'})

class RegistrationTests(APITestCase):
    def test_successful_registration(self):
        """Test registration creates user with hashed password and default student profile"""
        response = self.client.post('/api/auth/register/', {
            'username': 'john_doe',
            'email': 'john@example.com',
            'password': 'Password123!',
        }, format='json')

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['username'], 'john_doe')
        self.assertEqual(response.data['email'], 'john@example.com')
        self.assertEqual(response.data['role'], 'student')

        # Verify password is encrypted in database
        user = User.objects.get(username='john_doe')
        self.assertTrue(user.check_password('Password123!'))
        self.assertNotEqual(user.password, 'Password123!')
        self.assertEqual(user.profile.role, 'student')

    def test_duplicate_username_fails(self):
        """Registration with already taken username must return 400 Bad Request"""
        User.objects.create_user(username='existing_user', email='first@example.com', password='Password123!')

        response = self.client.post('/api/auth/register/', {
            'username': 'existing_user',
            'email': 'second@example.com',
            'password': 'Password123!',
        }, format='json')

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('username', response.data)

    def test_duplicate_email_fails(self):
        """Registration with already registered email must return 400 Bad Request"""
        User.objects.create_user(username='user1', email='duplicate@example.com', password='Password123!')

        response = self.client.post('/api/auth/register/', {
            'username': 'user2',
            'email': 'duplicate@example.com',
            'password': 'Password123!',
        }, format='json')

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('email', response.data)

    def test_short_password_fails(self):
        """Registration with password shorter than 8 chars must return 400 Bad Request"""
        response = self.client.post('/api/auth/register/', {
            'username': 'shortpass_user',
            'email': 'short@example.com',
            'password': '123',
        }, format='json')

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('password', response.data)

    def test_missing_required_fields_fails(self):
        """Registration without required fields must return 400 Bad Request"""
        response = self.client.post('/api/auth/register/', {
            'username': 'incomplete_user',
        }, format='json')

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)


class LoginAndJwtTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username='testuser',
            email='test@example.com',
            password='ValidPassword123!'
        )

    def test_login_returns_jwt_tokens(self):
        """Valid credentials return access and refresh JWT tokens"""
        response = self.client.post('/api/auth/login/', {
            'username': 'testuser',
            'password': 'ValidPassword123!'
        }, format='json')

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('access', response.data)
        self.assertIn('refresh', response.data)

    def test_login_incorrect_password_fails(self):
        """Incorrect password returns 401 Unauthorized"""
        response = self.client.post('/api/auth/login/', {
            'username': 'testuser',
            'password': 'WrongPassword999!'
        }, format='json')

        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_login_nonexistent_user_fails(self):
        """Non-existent username returns 401 Unauthorized"""
        response = self.client.post('/api/auth/login/', {
            'username': 'ghost_user',
            'password': 'ValidPassword123!'
        }, format='json')

        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_token_refresh_works(self):
        """Valid refresh token generates a new access token"""
        refresh = RefreshToken.for_user(self.user)
        response = self.client.post('/api/auth/token/refresh/', {
            'refresh': str(refresh)
        }, format='json')

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('access', response.data)

    def test_invalid_token_refresh_fails(self):
        """Invalid refresh token returns 401 Unauthorized"""
        response = self.client.post('/api/auth/token/refresh/', {
            'refresh': 'invalid-refresh-token'
        }, format='json')

        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)


class ProtectedEndpointAndDataIsolationTests(APITestCase):
    def setUp(self):
        self.user_a = User.objects.create_user(
            username='user_a',
            email='a@example.com',
            password='Password123!',
        )
        self.user_b = User.objects.create_user(
            username='user_b',
            email='b@example.com',
            password='Password123!',
        )

        self.expense_a = Expense.objects.create(
            user=self.user_a,
            category='Food',
            description='Lunch',
            amount='15.50',
            date=date(2026, 8, 28),
        )
        self.expense_b = Expense.objects.create(
            user=self.user_b,
            category='Travel',
            description='Train ticket',
            amount='25.00',
            date=date(2026, 8, 28),
        )

    def authenticate(self, user):
        token = RefreshToken.for_user(user).access_token
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

    def test_unauthenticated_request_is_rejected(self):
        """Endpoints reject unauthenticated requests with 401"""
        endpoints = ['/api/auth/me/', '/api/expenses/', '/api/incomes/', '/api/budgets/']
        for url in endpoints:
            response = self.client.get(url)
            self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED, f"Failed for {url}")

    def test_invalid_token_is_rejected(self):
        """Endpoints reject requests with an invalid Bearer token"""
        self.client.credentials(HTTP_AUTHORIZATION='Bearer invalid.fake.token')
        response = self.client.get('/api/auth/me/')
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_expired_token_is_rejected(self):
        """Endpoints reject requests with an expired Bearer token"""
        token = RefreshToken.for_user(self.user_a).access_token
        token['exp'] = 0
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')
        response = self.client.get('/api/auth/me/')
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_authenticated_profile_endpoint(self):
        """Authenticated user can fetch their own profile details"""
        self.authenticate(self.user_a)
        response = self.client.get('/api/auth/me/')

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['username'], 'user_a')
        self.assertEqual(response.data['email'], 'a@example.com')
        self.assertEqual(response.data['role'], 'student')

    def test_user_data_isolation(self):
        """User A can only view their own expenses and not User B's"""
        self.authenticate(self.user_a)
        response = self.client.get('/api/expenses/')

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        expense_ids = [item['id'] for item in response.data]
        self.assertIn(self.expense_a.id, expense_ids)
        self.assertNotIn(self.expense_b.id, expense_ids)

    def test_create_expense_assigns_authenticated_user(self):
        """Created expense is automatically associated with the requesting user"""
        self.authenticate(self.user_a)
        response = self.client.post('/api/expenses/', {
            'category': 'Education',
            'description': 'Books',
            'amount': '45.00',
            'date': '2026-08-28',
            'user': self.user_b.id  # attempt to spoof user_b
        }, format='json')

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        new_expense = Expense.objects.get(id=response.data['id'])
        self.assertEqual(new_expense.user, self.user_a)


class ExpenseManagementTests(APITestCase):
    def setUp(self):
        self.user_a = User.objects.create_user(
            username='student_alice',
            email='alice@example.com',
            password='Password123!'
        )
        self.user_b = User.objects.create_user(
            username='student_bob',
            email='bob@example.com',
            password='Password123!'
        )

        self.token_a = RefreshToken.for_user(self.user_a).access_token
        self.token_b = RefreshToken.for_user(self.user_b).access_token

        # Create baseline expenses for alice
        self.exp_a1 = Expense.objects.create(
            user=self.user_a,
            category='Food',
            description='Breakfast',
            amount='8.50',
            date=date(2026, 9, 1)
        )
        self.exp_a2 = Expense.objects.create(
            user=self.user_a,
            category='Travel',
            description='Metro ticket',
            amount='3.00',
            date=date(2026, 9, 2)
        )

        # Create baseline expense for bob
        self.exp_b1 = Expense.objects.create(
            user=self.user_b,
            category='Shopping',
            description='Backpack',
            amount='50.00',
            date=date(2026, 9, 1)
        )

    def auth_alice(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token_a}')

    def auth_bob(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token_b}')

    # --- Task 1: Expense Creation ---
    def test_create_expense_success(self):
        """Authenticated user can add an expense with valid fields and ownership association"""
        self.auth_alice()
        data = {
            'category': 'Education',
            'description': 'Textbook',
            'amount': '45.00',
            'date': '2026-09-02'
        }
        response = self.client.post('/api/expenses/', data, format='json')

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['category'], 'Education')
        self.assertEqual(response.data['amount'], '45.00')
        self.assertEqual(response.data['date'], '2026-09-02')

        created_exp = Expense.objects.get(id=response.data['id'])
        self.assertEqual(created_exp.user, self.user_a)
        self.assertEqual(created_exp.description, 'Textbook')

    def test_create_expense_validation_error(self):
        """Expense creation with invalid amount or missing fields fails with 400"""
        self.auth_alice()
        # Negative amount
        response = self.client.post('/api/expenses/', {
            'category': 'Food',
            'amount': '-10.00',
            'date': '2026-09-02'
        }, format='json')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

        # Missing date
        response = self.client.post('/api/expenses/', {
            'category': 'Food',
            'amount': '10.00'
        }, format='json')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    # --- Task 2: Expense Retrieval & Daily History Ordering ---
    def test_expense_list_retrieval_order(self):
        """User retrieves own expenses ordered chronologically by date descending"""
        self.auth_alice()
        response = self.client.get('/api/expenses/')

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data), 2)
        # 2026-09-02 (exp_a2) should be before 2026-09-01 (exp_a1)
        self.assertEqual(response.data[0]['id'], self.exp_a2.id)
        self.assertEqual(response.data[1]['id'], self.exp_a1.id)

    def test_expense_detail_retrieval(self):
        """User can retrieve own expense detail, but cannot retrieve another user's expense"""
        self.auth_alice()
        # Alice retrieves her own expense
        res_own = self.client.get(f'/api/expenses/{self.exp_a1.id}/')
        self.assertEqual(res_own.status_code, status.HTTP_200_OK)
        self.assertEqual(res_own.data['id'], self.exp_a1.id)

        # Alice tries to retrieve Bob's expense -> 404
        res_other = self.client.get(f'/api/expenses/{self.exp_b1.id}/')
        self.assertEqual(res_other.status_code, status.HTTP_404_NOT_FOUND)

    # --- Task 3: Expense Update ---
    def test_expense_update_put_and_patch(self):
        """User can update their own expense via PUT and PATCH"""
        self.auth_alice()

        # Full update (PUT)
        put_data = {
            'category': 'Food',
            'description': 'Fancy Dinner',
            'amount': '35.00',
            'date': '2026-09-01'
        }
        res_put = self.client.put(f'/api/expenses/{self.exp_a1.id}/', put_data, format='json')
        self.assertEqual(res_put.status_code, status.HTTP_200_OK)
        self.exp_a1.refresh_from_db()
        self.assertEqual(self.exp_a1.description, 'Fancy Dinner')
        self.assertEqual(str(self.exp_a1.amount), '35.00')

        # Partial update (PATCH)
        patch_data = {'amount': '40.00'}
        res_patch = self.client.patch(f'/api/expenses/{self.exp_a1.id}/', patch_data, format='json')
        self.assertEqual(res_patch.status_code, status.HTTP_200_OK)
        self.exp_a1.refresh_from_db()
        self.assertEqual(str(self.exp_a1.amount), '40.00')

    def test_expense_update_other_user_fails(self):
        """User cannot update another user's expense"""
        self.auth_alice()
        res = self.client.put(f'/api/expenses/{self.exp_b1.id}/', {
            'category': 'Shopping',
            'description': 'Hacked',
            'amount': '1.00',
            'date': '2026-09-01'
        }, format='json')
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)

        # Verify Bob's expense was not modified
        self.exp_b1.refresh_from_db()
        self.assertEqual(self.exp_b1.description, 'Backpack')

    def test_expense_update_invalid_data_fails(self):
        """Updating with invalid category or negative amount returns 400"""
        self.auth_alice()
        res = self.client.patch(f'/api/expenses/{self.exp_a1.id}/', {
            'amount': '-50.00'
        }, format='json')
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)

    # --- Task 4: Expense Deletion ---
    def test_expense_deletion_success(self):
        """User can delete their own expense"""
        self.auth_alice()
        res = self.client.delete(f'/api/expenses/{self.exp_a1.id}/')
        self.assertEqual(res.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(Expense.objects.filter(id=self.exp_a1.id).exists())

    def test_expense_deletion_other_user_fails(self):
        """User cannot delete another user's expense"""
        self.auth_alice()
        res = self.client.delete(f'/api/expenses/{self.exp_b1.id}/')
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)
        self.assertTrue(Expense.objects.filter(id=self.exp_b1.id).exists())

    # --- Task 5: Required Categorization ---
    def test_all_required_categories_accepted(self):
        """All 6 required categories (Food, Travel, Shopping, Education, Entertainment, Miscellaneous) work"""
        self.auth_alice()
        categories = ['Food', 'Travel', 'Shopping', 'Education', 'Entertainment', 'Miscellaneous']

        for cat in categories:
            res = self.client.post('/api/expenses/', {
                'category': cat,
                'description': f'Test {cat}',
                'amount': '15.00',
                'date': '2026-09-02'
            }, format='json')
            self.assertEqual(res.status_code, status.HTTP_201_CREATED, f"Failed for category: {cat}")
            self.assertEqual(res.data['category'], cat)

    def test_invalid_category_rejected(self):
        """Unsupported category is rejected with HTTP 400"""
        self.auth_alice()
        res = self.client.post('/api/expenses/', {
            'category': 'Gambling',
            'description': 'Casino',
            'amount': '100.00',
            'date': '2026-09-02'
        }, format='json')
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('category', res.data)

    def test_missing_category_rejected(self):
        """Expense creation with missing category fails with 400"""
        self.auth_alice()
        res = self.client.post('/api/expenses/', {
            'description': 'No category',
            'amount': '25.00',
            'date': '2026-09-02'
        }, format='json')
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('category', res.data)

    def test_update_expense_category(self):
        """Updating expense category changes category and validates allowed choices"""
        self.auth_alice()
        res = self.client.patch(f'/api/expenses/{self.exp_a1.id}/', {
            'category': 'Entertainment'
        }, format='json')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.exp_a1.refresh_from_db()
        self.assertEqual(self.exp_a1.category, 'Entertainment')

        # Invalid category update fails
        res_invalid = self.client.patch(f'/api/expenses/{self.exp_a1.id}/', {
            'category': 'NotAValidCategory'
        }, format='json')
        self.assertEqual(res_invalid.status_code, status.HTTP_400_BAD_REQUEST)


class IncomeManagementTests(APITestCase):
    def setUp(self):
        self.user_a = User.objects.create_user(
            username='student_claire',
            email='claire@example.com',
            password='Password123!'
        )
        self.user_b = User.objects.create_user(
            username='student_david',
            email='david@example.com',
            password='Password123!'
        )

        self.token_a = RefreshToken.for_user(self.user_a).access_token
        self.token_b = RefreshToken.for_user(self.user_b).access_token

        # Create baseline income for claire
        self.inc_a1 = Income.objects.create(
            user=self.user_a,
            source='Pocket Money',
            description='Monthly allowance',
            amount='100.00',
            date=date(2026, 9, 1)
        )
        self.inc_a2 = Income.objects.create(
            user=self.user_a,
            source='Scholarship',
            description='Merit grant',
            amount='250.00',
            date=date(2026, 9, 3)
        )

        # Create baseline income for david
        self.inc_b1 = Income.objects.create(
            user=self.user_b,
            source='Freelance Income',
            description='Web design',
            amount='300.00',
            date=date(2026, 9, 2)
        )

    def auth_claire(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token_a}')

    def auth_david(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token_b}')

    # --- Task 1: Add Income & Required Sources ---
    def test_create_income_success(self):
        """Authenticated user can add an income record with valid data and user association"""
        self.auth_claire()
        data = {
            'source': 'Freelance Income',
            'description': 'Tutoring session',
            'amount': '45.00',
            'date': '2026-09-04'
        }
        res = self.client.post('/api/incomes/', data, format='json')

        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertEqual(res.data['source'], 'Freelance Income')
        self.assertEqual(res.data['amount'], '45.00')
        self.assertEqual(res.data['date'], '2026-09-04')

        created_inc = Income.objects.get(id=res.data['id'])
        self.assertEqual(created_inc.user, self.user_a)
        self.assertEqual(created_inc.description, 'Tutoring session')

    def test_required_income_sources_handled(self):
        """Specified sources (Pocket Money, Scholarship, Freelance Income) are all accepted"""
        self.auth_claire()
        sources = ['Pocket Money', 'Scholarship', 'Freelance Income']

        for src in sources:
            res = self.client.post('/api/incomes/', {
                'source': src,
                'description': f'Test {src}',
                'amount': '50.00',
                'date': '2026-09-04'
            }, format='json')
            self.assertEqual(res.status_code, status.HTTP_201_CREATED, f"Failed for source: {src}")
            self.assertEqual(res.data['source'], src)

    def test_create_income_validation_error(self):
        """Invalid amount (<= 0) or missing required fields returns 400 Bad Request"""
        self.auth_claire()

        # Negative amount
        res = self.client.post('/api/incomes/', {
            'source': 'Pocket Money',
            'amount': '-20.00',
            'date': '2026-09-04'
        }, format='json')
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)

        # Missing date
        res2 = self.client.post('/api/incomes/', {
            'source': 'Pocket Money',
            'amount': '20.00'
        }, format='json')
        self.assertEqual(res2.status_code, status.HTTP_400_BAD_REQUEST)

    # --- Task 2: View Income History ---
    def test_income_list_retrieval_order(self):
        """User retrieves own income records in descending chronological order"""
        self.auth_claire()
        res = self.client.get('/api/incomes/')

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(len(res.data), 2)
        # 2026-09-03 (inc_a2) should be before 2026-09-01 (inc_a1)
        self.assertEqual(res.data[0]['id'], self.inc_a2.id)
        self.assertEqual(res.data[1]['id'], self.inc_a1.id)

    def test_income_detail_retrieval(self):
        """User can view own income detail, but cannot view another user's income"""
        self.auth_claire()
        # Claire views own income
        res_own = self.client.get(f'/api/incomes/{self.inc_a1.id}/')
        self.assertEqual(res_own.status_code, status.HTTP_200_OK)
        self.assertEqual(res_own.data['id'], self.inc_a1.id)

        # Claire tries to view David's income -> 404
        res_other = self.client.get(f'/api/incomes/{self.inc_b1.id}/')
        self.assertEqual(res_other.status_code, status.HTTP_404_NOT_FOUND)

    # --- Task 3: Manage Income Records (Update & Delete) ---
    def test_income_update_put_and_patch(self):
        """User can update own income record"""
        self.auth_claire()

        # PUT
        res_put = self.client.put(f'/api/incomes/{self.inc_a1.id}/', {
            'source': 'Pocket Money',
            'description': 'Updated allowance',
            'amount': '120.00',
            'date': '2026-09-01'
        }, format='json')
        self.assertEqual(res_put.status_code, status.HTTP_200_OK)
        self.inc_a1.refresh_from_db()
        self.assertEqual(self.inc_a1.description, 'Updated allowance')
        self.assertEqual(str(self.inc_a1.amount), '120.00')

        # PATCH
        res_patch = self.client.patch(f'/api/incomes/{self.inc_a1.id}/', {
            'amount': '130.00'
        }, format='json')
        self.assertEqual(res_patch.status_code, status.HTTP_200_OK)
        self.inc_a1.refresh_from_db()
        self.assertEqual(str(self.inc_a1.amount), '130.00')

    def test_income_update_other_user_fails(self):
        """User cannot update another user's income record"""
        self.auth_claire()
        res = self.client.put(f'/api/incomes/{self.inc_b1.id}/', {
            'source': 'Freelance Income',
            'description': 'Hacked income',
            'amount': '1.00',
            'date': '2026-09-02'
        }, format='json')
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)

        self.inc_b1.refresh_from_db()
        self.assertEqual(self.inc_b1.description, 'Web design')

    def test_income_deletion_success(self):
        """User can delete own income record"""
        self.auth_claire()
        res = self.client.delete(f'/api/incomes/{self.inc_a1.id}/')
        self.assertEqual(res.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(Income.objects.filter(id=self.inc_a1.id).exists())

    def test_income_deletion_other_user_fails(self):
        """User cannot delete another user's income record"""
        self.auth_claire()
        res = self.client.delete(f'/api/incomes/{self.inc_b1.id}/')
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)
        self.assertTrue(Income.objects.filter(id=self.inc_b1.id).exists())


class DashboardSummaryTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username='summary_user',
            email='summary@example.com',
            password='Password123!'
        )
        self.other_user = User.objects.create_user(
            username='other_user',
            email='other@example.com',
            password='Password123!'
        )
        self.token = RefreshToken.for_user(self.user).access_token

        # Incomes for current user: 200 + 300 = 500
        Income.objects.create(
            user=self.user,
            source='Pocket Money',
            amount='200.00',
            date=date(2026, 9, 1)
        )
        Income.objects.create(
            user=self.user,
            source='Scholarship',
            amount='300.00',
            date=date(2026, 9, 3)
        )

        # Expenses for current user: 50 + 150 = 200
        Expense.objects.create(
            user=self.user,
            category='Food',
            amount='50.00',
            date=date(2026, 9, 2)
        )
        Expense.objects.create(
            user=self.user,
            category='Education',
            amount='150.00',
            date=date(2026, 9, 4)
        )

        # Other user's data (must not affect summary)
        Income.objects.create(
            user=self.other_user,
            source='Salary',
            amount='1000.00',
            date=date(2026, 9, 1)
        )
        Expense.objects.create(
            user=self.other_user,
            category='Travel',
            amount='400.00',
            date=date(2026, 9, 1)
        )

    def test_dashboard_summary_calculation(self):
        """Dashboard summary computes Total Income, Total Expenses, and Remaining Amount correctly"""
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token}')
        res = self.client.get('/api/dashboard/summary/')

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data['total_income'], '500.00')
        self.assertEqual(res.data['total_expenses'], '200.00')
        self.assertEqual(res.data['remaining_amount'], '300.00')
        self.assertEqual(res.data['income_count'], 2)
        self.assertEqual(res.data['expense_count'], 2)

    def test_combined_recent_transactions_order(self):
        """Combined recent transactions list contains both incomes & expenses chronologically sorted"""
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token}')
        res = self.client.get('/api/dashboard/summary/')

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        transactions = res.data['recent_transactions']
        self.assertEqual(len(transactions), 4)

        # Dates should be strictly descending: 2026-09-04 (exp), 2026-09-03 (inc), 2026-09-02 (exp), 2026-09-01 (inc)
        dates = [t['date'] for t in transactions]
        self.assertEqual(dates, sorted(dates, reverse=True))

        types = [t['type'] for t in transactions]
        self.assertIn('income', types)
        self.assertIn('expense', types)


class BudgetManagementTests(APITestCase):
    def setUp(self):
        self.user_a = User.objects.create_user(
            username='student_emma',
            email='emma@example.com',
            password='Password123!'
        )
        self.user_b = User.objects.create_user(
            username='student_frank',
            email='frank@example.com',
            password='Password123!'
        )

        self.token_a = RefreshToken.for_user(self.user_a).access_token
        self.token_b = RefreshToken.for_user(self.user_b).access_token

        # Baseline budget for emma
        self.b_a1 = Budget.objects.create(
            user=self.user_a,
            category='Food',
            amount='300.00',
            month=date(2026, 9, 1)
        )
        self.b_a2 = Budget.objects.create(
            user=self.user_a,
            category='Travel',
            amount='100.00',
            month=date(2026, 9, 1)
        )

        # Baseline budget for frank
        self.b_b1 = Budget.objects.create(
            user=self.user_b,
            category='Food',
            amount='400.00',
            month=date(2026, 9, 1)
        )

        # Some expenses for Emma in September
        Expense.objects.create(
            user=self.user_a,
            category='Food',
            amount='75.00',
            date=date(2026, 9, 2)
        )
        Expense.objects.create(
            user=self.user_a,
            category='Food',
            amount='50.00',
            date=date(2026, 9, 3)
        )

    def auth_emma(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token_a}')

    def auth_frank(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token_b}')

    def test_create_budget_success(self):
        """Authenticated user can allocate a monthly budget for a category"""
        self.auth_emma()
        data = {
            'category': 'Education',
            'amount': '150.00',
            'month': '2026-09-01'
        }
        res = self.client.post('/api/budgets/', data, format='json')

        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertEqual(res.data['category'], 'Education')
        self.assertEqual(res.data['amount'], '150.00')

        b = Budget.objects.get(id=res.data['id'])
        self.assertEqual(b.user, self.user_a)

    def test_create_budget_all_required_categories(self):
        """All 6 predefined categories (Food, Travel, Shopping, Education, Entertainment, Miscellaneous) work for budget"""
        self.auth_emma()
        categories = ['Food', 'Travel', 'Shopping', 'Education', 'Entertainment', 'Miscellaneous']

        for cat in categories:
            res = self.client.post('/api/budgets/', {
                'category': cat,
                'amount': '100.00',
                'month': '2026-10-01'
            }, format='json')
            self.assertEqual(res.status_code, status.HTTP_201_CREATED, f"Failed for category: {cat}")
            self.assertEqual(res.data['category'], cat)

    def test_create_budget_validation_error(self):
        """Negative amount, invalid category, or missing month returns 400 Bad Request"""
        self.auth_emma()

        # Negative amount
        res = self.client.post('/api/budgets/', {
            'category': 'Food',
            'amount': '-50.00',
            'month': '2026-09-01'
        }, format='json')
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)

        # Invalid category
        res2 = self.client.post('/api/budgets/', {
            'category': 'InvalidCategory',
            'amount': '50.00',
            'month': '2026-09-01'
        }, format='json')
        self.assertEqual(res2.status_code, status.HTTP_400_BAD_REQUEST)

        # Missing category
        res3 = self.client.post('/api/budgets/', {
            'amount': '50.00',
            'month': '2026-09-01'
        }, format='json')
        self.assertEqual(res3.status_code, status.HTTP_400_BAD_REQUEST)

        # Missing amount
        res4 = self.client.post('/api/budgets/', {
            'category': 'Food',
            'month': '2026-09-01'
        }, format='json')
        self.assertEqual(res4.status_code, status.HTTP_400_BAD_REQUEST)

    def test_budget_list_retrieval_and_month_filter(self):
        """User retrieves their own budgets with optional month filter"""
        self.auth_emma()
        res = self.client.get('/api/budgets/?month=2026-09-01')

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(len(res.data), 2)
        categories = [b['category'] for b in res.data]
        self.assertIn('Food', categories)
        self.assertIn('Travel', categories)

    def test_budget_detail_retrieval(self):
        """User can view own budget detail, but cannot view another user's budget"""
        self.auth_emma()
        # Emma views own budget
        res_own = self.client.get(f'/api/budgets/{self.b_a1.id}/')
        self.assertEqual(res_own.status_code, status.HTTP_200_OK)
        self.assertEqual(res_own.data['id'], self.b_a1.id)

        # Emma tries to view Frank's budget -> 404
        res_other = self.client.get(f'/api/budgets/{self.b_b1.id}/')
        self.assertEqual(res_other.status_code, status.HTTP_404_NOT_FOUND)

    def test_budget_update_put_and_patch(self):
        """User can update own budget allocation"""
        self.auth_emma()

        # PUT
        res_put = self.client.put(f'/api/budgets/{self.b_a1.id}/', {
            'category': 'Food',
            'amount': '350.00',
            'month': '2026-09-01'
        }, format='json')
        self.assertEqual(res_put.status_code, status.HTTP_200_OK)
        self.b_a1.refresh_from_db()
        self.assertEqual(str(self.b_a1.amount), '350.00')

        # PATCH
        res_patch = self.client.patch(f'/api/budgets/{self.b_a1.id}/', {
            'amount': '375.00'
        }, format='json')
        self.assertEqual(res_patch.status_code, status.HTTP_200_OK)
        self.b_a1.refresh_from_db()
        self.assertEqual(str(self.b_a1.amount), '375.00')

    def test_budget_update_other_user_fails(self):
        """User cannot update another user's budget allocation"""
        self.auth_emma()
        res = self.client.put(f'/api/budgets/{self.b_b1.id}/', {
            'category': 'Food',
            'amount': '1.00',
            'month': '2026-09-01'
        }, format='json')
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)

        self.b_b1.refresh_from_db()
        self.assertEqual(str(self.b_b1.amount), '400.00')

    def test_budget_deletion_success(self):
        """User can delete own budget allocation"""
        self.auth_emma()
        res = self.client.delete(f'/api/budgets/{self.b_a1.id}/')
        self.assertEqual(res.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(Budget.objects.filter(id=self.b_a1.id).exists())

    def test_budget_deletion_other_user_fails(self):
        """User cannot delete another user's budget allocation"""
        self.auth_emma()
        res = self.client.delete(f'/api/budgets/{self.b_b1.id}/')
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)
        self.assertTrue(Budget.objects.filter(id=self.b_b1.id).exists())

    def test_budget_summary_calculation(self):
        """Budget summary computes allocated vs actual spent for each category in the selected month"""
        self.auth_emma()
        res = self.client.get('/api/budgets/summary/?month=2026-09-01')

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data['total_budget'], '400.00') # 300 (Food) + 100 (Travel)
        self.assertEqual(res.data['total_spent'], '125.00') # 75 + 50 (Food)
        self.assertEqual(res.data['total_remaining'], '275.00')

        # Check Food category detail
        food_cat = next(c for c in res.data['categories'] if c['category'] == 'Food')
        self.assertEqual(food_cat['allocated_budget'], '300.00')
        self.assertEqual(food_cat['actual_spent'], '125.00')
        self.assertEqual(food_cat['remaining_budget'], '175.00')
        self.assertAlmostEqual(food_cat['utilization_percentage'], 41.7, places=1)
        self.assertFalse(food_cat['is_over_budget'])


class Milestone2IntegrationFlowTests(APITestCase):
    """
    End-to-End Integration, Testing & Validation for Milestone 2:
    - Incomes, Expenses, Categorization, Budgeting & Dashboard Summary Integration
    - Complete Lifecycle Flow: Add Transactions -> Dashboard Sync -> Category Budgets -> Utilization & Overspending -> Live Updates
    - Multi-User Data Isolation across all 5 Milestone 2 endpoints
    - Cross-Module Validation & Error Handling
    """

    def setUp(self):
        self.primary_user = User.objects.create_user(
            username='integration_lead',
            email='lead@example.com',
            password='Password123!'
        )
        self.secondary_user = User.objects.create_user(
            username='other_student',
            email='other@example.com',
            password='Password123!'
        )
        self.primary_token = RefreshToken.for_user(self.primary_user).access_token
        self.secondary_token = RefreshToken.for_user(self.secondary_user).access_token

    def auth_primary(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.primary_token}')

    def auth_secondary(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.secondary_token}')

    def test_complete_milestone2_lifecycle_flow(self):
        """
        Full End-to-End Transaction & Budget Lifecycle:
        1. Record incomes across sources
        2. Record expenses across all 6 project categories
        3. Verify Dashboard Financial Summary & Combined Chronological Activity
        4. Create Category Budgets & verify Budget Summary calculations with Overspending detection
        5. Update and Delete transactions -> verify live dashboard & budget synchronization
        """
        self.auth_primary()

        # Step 1: Record 3 Incomes across different sources
        inc1 = self.client.post('/api/incomes/', {
            'source': 'Pocket Money',
            'amount': '500.00',
            'date': '2026-09-01',
            'description': 'Monthly allowance from parents'
        }, format='json')
        self.assertEqual(inc1.status_code, status.HTTP_201_CREATED)

        inc2 = self.client.post('/api/incomes/', {
            'source': 'Scholarship',
            'amount': '1000.00',
            'date': '2026-09-02',
            'description': 'Academic semester grant'
        }, format='json')
        self.assertEqual(inc2.status_code, status.HTTP_201_CREATED)

        inc3 = self.client.post('/api/incomes/', {
            'source': 'Freelance Income',
            'amount': '300.00',
            'date': '2026-09-03',
            'description': 'Graphic design project'
        }, format='json')
        self.assertEqual(inc3.status_code, status.HTTP_201_CREATED)
        # Total Incomes = 500 + 1000 + 300 = 1800.00

        # Step 2: Record 6 Expenses across all 6 predefined project categories
        expenses_data = [
            ('Food', '120.00', '2026-09-01', 'Groceries & Dining'),
            ('Travel', '45.00', '2026-09-02', 'Bus pass & Train tickets'),
            ('Shopping', '85.00', '2026-09-02', 'Clothing & Shoes'),
            ('Education', '200.00', '2026-09-03', 'Course Textbooks'),
            ('Entertainment', '50.00', '2026-09-04', 'Concert ticket'),
            ('Miscellaneous', '30.00', '2026-09-04', 'Stationery & Supplies')
        ]
        created_exp_ids = {}
        for cat, amt, dt, desc in expenses_data:
            res = self.client.post('/api/expenses/', {
                'category': cat,
                'amount': amt,
                'date': dt,
                'description': desc
            }, format='json')
            self.assertEqual(res.status_code, status.HTTP_201_CREATED)
            created_exp_ids[cat] = res.data['id']
        # Total Expenses = 120 + 45 + 85 + 200 + 50 + 30 = 530.00

        # Step 3: Check Dashboard Financial Summary Integration
        dash_res = self.client.get('/api/dashboard/summary/')
        self.assertEqual(dash_res.status_code, status.HTTP_200_OK)
        self.assertEqual(dash_res.data['total_income'], '1800.00')
        self.assertEqual(dash_res.data['total_expenses'], '530.00')
        self.assertEqual(dash_res.data['remaining_amount'], '1270.00')
        self.assertEqual(dash_res.data['income_count'], 3)
        self.assertEqual(dash_res.data['expense_count'], 6)

        # Verify combined chronological transaction activity feed (9 items total)
        recent_txs = dash_res.data['recent_transactions']
        self.assertEqual(len(recent_txs), 9)
        dates = [tx['date'] for tx in recent_txs]
        self.assertEqual(dates, sorted(dates, reverse=True))

        # Step 4: Create Category Budgets for 2026-09-01
        # Food: ₹100.00 budget (Spent ₹120.00 -> should trigger is_over_budget: True)
        b_food = self.client.post('/api/budgets/', {
            'category': 'Food',
            'amount': '100.00',
            'month': '2026-09-01'
        }, format='json')
        self.assertEqual(b_food.status_code, status.HTTP_201_CREATED)

        # Travel: ₹100.00 budget (Spent ₹45.00 -> 45.0% utilization)
        b_travel = self.client.post('/api/budgets/', {
            'category': 'Travel',
            'amount': '100.00',
            'month': '2026-09-01'
        }, format='json')
        self.assertEqual(b_travel.status_code, status.HTTP_201_CREATED)

        # Education: ₹200.00 budget (Spent ₹200.00 -> 100.0% utilization)
        b_edu = self.client.post('/api/budgets/', {
            'category': 'Education',
            'amount': '200.00',
            'month': '2026-09-01'
        }, format='json')
        self.assertEqual(b_edu.status_code, status.HTTP_201_CREATED)
        # Total Budget = 100 + 100 + 200 = 400.00

        # Step 5: Check Budget Summary Analytics Integration
        b_sum_res = self.client.get('/api/budgets/summary/?month=2026-09-01')
        self.assertEqual(b_sum_res.status_code, status.HTTP_200_OK)
        self.assertEqual(b_sum_res.data['total_budget'], '400.00')
        self.assertEqual(b_sum_res.data['total_spent'], '530.00')
        self.assertEqual(b_sum_res.data['total_remaining'], '-130.00')

        cats = {c['category']: c for c in b_sum_res.data['categories']}

        # Food Check (Over budget)
        self.assertEqual(cats['Food']['allocated_budget'], '100.00')
        self.assertEqual(cats['Food']['actual_spent'], '120.00')
        self.assertEqual(cats['Food']['remaining_budget'], '-20.00')
        self.assertEqual(cats['Food']['utilization_percentage'], 120.0)
        self.assertTrue(cats['Food']['is_over_budget'])

        # Travel Check (Within budget)
        self.assertEqual(cats['Travel']['allocated_budget'], '100.00')
        self.assertEqual(cats['Travel']['actual_spent'], '45.00')
        self.assertEqual(cats['Travel']['remaining_budget'], '55.00')
        self.assertEqual(cats['Travel']['utilization_percentage'], 45.0)
        self.assertFalse(cats['Travel']['is_over_budget'])

        # Education Check (Exact 100% budget)
        self.assertEqual(cats['Education']['allocated_budget'], '200.00')
        self.assertEqual(cats['Education']['actual_spent'], '200.00')
        self.assertEqual(cats['Education']['remaining_budget'], '0.00')
        self.assertEqual(cats['Education']['utilization_percentage'], 100.0)
        self.assertFalse(cats['Education']['is_over_budget'])

        # Step 6: Update Food expense from ₹120.00 to ₹80.00 -> check dynamic recalculations
        food_exp_id = created_exp_ids['Food']
        update_exp_res = self.client.patch(f'/api/expenses/{food_exp_id}/', {
            'amount': '80.00'
        }, format='json')
        self.assertEqual(update_exp_res.status_code, status.HTTP_200_OK)

        # Verify Dashboard reflects updated expense: Total Expenses = 530 - 40 = 490.00
        dash_res_updated = self.client.get('/api/dashboard/summary/')
        self.assertEqual(dash_res_updated.data['total_expenses'], '490.00')
        self.assertEqual(dash_res_updated.data['remaining_amount'], '1310.00')

        # Verify Budget Summary reflects Food is no longer over budget (80/100 = 80%)
        b_sum_updated = self.client.get('/api/budgets/summary/?month=2026-09-01')
        cats_updated = {c['category']: c for c in b_sum_updated.data['categories']}
        self.assertEqual(cats_updated['Food']['actual_spent'], '80.00')
        self.assertEqual(cats_updated['Food']['remaining_budget'], '20.00')
        self.assertEqual(cats_updated['Food']['utilization_percentage'], 80.0)
        self.assertFalse(cats_updated['Food']['is_over_budget'])

        # Step 7: Delete Freelance Income (₹300.00) -> check dynamic recalculation
        inc3_id = inc3.data['id']
        del_inc_res = self.client.delete(f'/api/incomes/{inc3_id}/')
        self.assertEqual(del_inc_res.status_code, status.HTTP_204_NO_CONTENT)

        # Verify Dashboard reflects deleted income: Total Income = 1800 - 300 = 1500.00
        dash_res_after_del = self.client.get('/api/dashboard/summary/')
        self.assertEqual(dash_res_after_del.data['total_income'], '1500.00')
        self.assertEqual(dash_res_after_del.data['remaining_amount'], '1010.00')
        self.assertEqual(dash_res_after_del.data['income_count'], 2)

    def test_multi_user_data_isolation_across_all_milestone2_apis(self):
        """
        Verify strict user separation across all 5 Milestone 2 endpoints:
        - Incomes, Expenses, Budgets, Dashboard Summary, Budget Summary
        - One user cannot view, edit, delete, or aggregate another user's financial data
        """
        # User 1 sets up data
        self.auth_primary()
        inc = self.client.post('/api/incomes/', {
            'source': 'Salary',
            'amount': '2500.00',
            'date': '2026-09-01'
        }, format='json')
        exp = self.client.post('/api/expenses/', {
            'category': 'Food',
            'amount': '150.00',
            'date': '2026-09-02'
        }, format='json')
        bgt = self.client.post('/api/budgets/', {
            'category': 'Food',
            'amount': '300.00',
            'month': '2026-09-01'
        }, format='json')

        inc_id = inc.data['id']
        exp_id = exp.data['id']
        bgt_id = bgt.data['id']

        # User 2 logs in with zero records
        self.auth_secondary()

        # 1. Incomes isolation
        res_inc = self.client.get('/api/incomes/')
        self.assertEqual(res_inc.data, [])
        self.assertEqual(self.client.get(f'/api/incomes/{inc_id}/').status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(self.client.put(f'/api/incomes/{inc_id}/', {'source': 'Pocket Money', 'amount': '10.00', 'date': '2026-09-01'}, format='json').status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(self.client.delete(f'/api/incomes/{inc_id}/').status_code, status.HTTP_404_NOT_FOUND)

        # 2. Expenses isolation
        res_exp = self.client.get('/api/expenses/')
        self.assertEqual(res_exp.data, [])
        self.assertEqual(self.client.get(f'/api/expenses/{exp_id}/').status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(self.client.put(f'/api/expenses/{exp_id}/', {'category': 'Food', 'amount': '10.00', 'date': '2026-09-01'}, format='json').status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(self.client.delete(f'/api/expenses/{exp_id}/').status_code, status.HTTP_404_NOT_FOUND)

        # 3. Budgets isolation
        res_bgt = self.client.get('/api/budgets/')
        self.assertEqual(res_bgt.data, [])
        self.assertEqual(self.client.get(f'/api/budgets/{bgt_id}/').status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(self.client.put(f'/api/budgets/{bgt_id}/', {'category': 'Food', 'amount': '10.00', 'month': '2026-09-01'}, format='json').status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(self.client.delete(f'/api/budgets/{bgt_id}/').status_code, status.HTTP_404_NOT_FOUND)

        # 4. Dashboard Summary isolation (must show 0.00 without leaking User 1 data)
        dash_res = self.client.get('/api/dashboard/summary/')
        self.assertEqual(dash_res.data['total_income'], '0.00')
        self.assertEqual(dash_res.data['total_expenses'], '0.00')
        self.assertEqual(dash_res.data['remaining_amount'], '0.00')
        self.assertEqual(dash_res.data['income_count'], 0)
        self.assertEqual(dash_res.data['expense_count'], 0)
        self.assertEqual(dash_res.data['recent_transactions'], [])

        # 5. Budget Summary isolation (must show 0.00 allocated and 0.00 spent)
        b_sum_res = self.client.get('/api/budgets/summary/?month=2026-09-01')
        self.assertEqual(b_sum_res.data['total_budget'], '0.00')
        self.assertEqual(b_sum_res.data['total_spent'], '0.00')
        self.assertEqual(b_sum_res.data['total_remaining'], '0.00')
        for c in b_sum_res.data['categories']:
            self.assertEqual(c['allocated_budget'], '0.00')
            self.assertEqual(c['actual_spent'], '0.00')
            self.assertEqual(c['remaining_budget'], '0.00')
            self.assertEqual(c['utilization_percentage'], 0.0)
            self.assertFalse(c['is_over_budget'])

    def test_validation_rules_across_all_milestone2_modules(self):
        """
        Validate data constraints across all Milestone 2 modules:
        - Rejection of negative & zero amounts (₹ <= 0)
        - Rejection of invalid/undefined categories & income sources
        - Rejection of missing required fields
        """
        self.auth_primary()

        # 1. Zero Amount Rejections
        self.assertEqual(self.client.post('/api/expenses/', {'category': 'Food', 'amount': '0.00', 'date': '2026-09-01'}, format='json').status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(self.client.post('/api/incomes/', {'source': 'Pocket Money', 'amount': '0.00', 'date': '2026-09-01'}, format='json').status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(self.client.post('/api/budgets/', {'category': 'Food', 'amount': '0.00', 'month': '2026-09-01'}, format='json').status_code, status.HTTP_400_BAD_REQUEST)

        # 2. Negative Amount Rejections
        self.assertEqual(self.client.post('/api/expenses/', {'category': 'Food', 'amount': '-25.00', 'date': '2026-09-01'}, format='json').status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(self.client.post('/api/incomes/', {'source': 'Pocket Money', 'amount': '-50.00', 'date': '2026-09-01'}, format='json').status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(self.client.post('/api/budgets/', {'category': 'Food', 'amount': '-100.00', 'month': '2026-09-01'}, format='json').status_code, status.HTTP_400_BAD_REQUEST)

        # 3. Invalid Choice Rejections
        self.assertEqual(self.client.post('/api/expenses/', {'category': 'InvalidCategory', 'amount': '25.00', 'date': '2026-09-01'}, format='json').status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(self.client.post('/api/incomes/', {'source': 'InvalidSource', 'amount': '50.00', 'date': '2026-09-01'}, format='json').status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(self.client.post('/api/budgets/', {'category': 'InvalidCategory', 'amount': '100.00', 'month': '2026-09-01'}, format='json').status_code, status.HTTP_400_BAD_REQUEST)


class Milestone2ComprehensiveIntegrationTests(APITestCase):
    """
    Detailed verification of:
    1. React -> API -> DB -> Response -> UI Data Flows for Expenses & Incomes
    2. Category Consistency across all 6 Predefined Categories
    3. Category Shifting Dynamic Recalculations across Budgets and Dashboards
    4. Non-Existent Record Handling (404s)
    5. Unauthenticated Request Rejection (401s)
    6. Cross-User Data Isolation
    7. Edge-case Invalid Inputs (Bad dates, invalid types, zero/negative amounts)
    8. Budget Upsert Logic
    """

    def setUp(self):
        self.user1 = User.objects.create_user(
            username='integ_user1',
            email='integ1@example.com',
            password='Password123!'
        )
        self.user2 = User.objects.create_user(
            username='integ_user2',
            email='integ2@example.com',
            password='Password123!'
        )
        self.token1 = RefreshToken.for_user(self.user1).access_token
        self.token2 = RefreshToken.for_user(self.user2).access_token

    def auth_user1(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token1}')

    def auth_user2(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token2}')

    def clear_auth(self):
        self.client.credentials()

    def test_react_flow_expense_to_db_and_response(self):
        """React -> Expense API -> DB -> Response flow verification"""
        self.auth_user1()
        payload = {
            'category': 'Education',
            'amount': '89.99',
            'date': '2026-09-05',
            'description': 'Online Course Subscription'
        }
        res = self.client.post('/api/expenses/', payload, format='json')
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertEqual(res.data['category'], 'Education')
        self.assertEqual(res.data['amount'], '89.99')
        self.assertEqual(res.data['date'], '2026-09-05')
        self.assertEqual(res.data['description'], 'Online Course Subscription')

        # Check DB persistence
        exp = Expense.objects.get(id=res.data['id'])
        self.assertEqual(exp.user, self.user1)
        self.assertEqual(str(exp.amount), '89.99')
        self.assertEqual(exp.category, 'Education')

    def test_react_flow_income_to_db_and_response(self):
        """React -> Income API -> DB -> Response flow verification"""
        self.auth_user1()
        payload = {
            'source': 'Scholarship',
            'amount': '750.00',
            'date': '2026-09-01',
            'description': 'Merit Award'
        }
        res = self.client.post('/api/incomes/', payload, format='json')
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertEqual(res.data['source'], 'Scholarship')
        self.assertEqual(res.data['amount'], '750.00')
        self.assertEqual(res.data['date'], '2026-09-01')
        self.assertEqual(res.data['description'], 'Merit Award')

        # Check DB persistence
        inc = Income.objects.get(id=res.data['id'])
        self.assertEqual(inc.user, self.user1)
        self.assertEqual(str(inc.amount), '750.00')
        self.assertEqual(inc.source, 'Scholarship')

    def test_expense_category_consistency_across_all_six_categories(self):
        """Verify all 6 categories are stored and retrieved consistently"""
        self.auth_user1()
        categories = ['Food', 'Travel', 'Shopping', 'Education', 'Entertainment', 'Miscellaneous']

        for idx, cat in enumerate(categories, start=1):
            res = self.client.post('/api/expenses/', {
                'category': cat,
                'amount': f'{idx * 10}.00',
                'date': '2026-09-01',
                'description': f'Testing {cat}'
            }, format='json')
            self.assertEqual(res.status_code, status.HTTP_201_CREATED)

        # Retrieve all expenses and verify category presence
        list_res = self.client.get('/api/expenses/')
        retrieved_cats = {item['category'] for item in list_res.data}
        self.assertEqual(retrieved_cats, set(categories))

    def test_expense_category_shift_updates_budget_and_dashboard(self):
        """
        Changing an expense category dynamically shifts actual spent amounts
        in the budget summary and updates the transaction feed.
        """
        self.auth_user1()
        # Set budgets
        self.client.post('/api/budgets/', {'category': 'Food', 'amount': '100.00', 'month': '2026-09-01'}, format='json')
        self.client.post('/api/budgets/', {'category': 'Entertainment', 'amount': '100.00', 'month': '2026-09-01'}, format='json')

        # Add Expense in Food
        exp_res = self.client.post('/api/expenses/', {
            'category': 'Food',
            'amount': '60.00',
            'date': '2026-09-02',
            'description': 'Dinner'
        }, format='json')
        exp_id = exp_res.data['id']

        # Verify initial budget: Food spent = 60.00, Entertainment spent = 0.00
        b_sum1 = self.client.get('/api/budgets/summary/?month=2026-09-01')
        cats1 = {c['category']: c for c in b_sum1.data['categories']}
        self.assertEqual(cats1['Food']['actual_spent'], '60.00')
        self.assertEqual(cats1['Entertainment']['actual_spent'], '0.00')

        # Shift category from Food to Entertainment
        patch_res = self.client.patch(f'/api/expenses/{exp_id}/', {'category': 'Entertainment'}, format='json')
        self.assertEqual(patch_res.status_code, status.HTTP_200_OK)

        # Verify budget dynamically recalculated: Food spent = 0.00, Entertainment spent = 60.00
        b_sum2 = self.client.get('/api/budgets/summary/?month=2026-09-01')
        cats2 = {c['category']: c for c in b_sum2.data['categories']}
        self.assertEqual(cats2['Food']['actual_spent'], '0.00')
        self.assertEqual(cats2['Entertainment']['actual_spent'], '60.00')

        # Verify dashboard summary remains consistent in total expenses (60.00) and reflects new category in recent transactions
        dash_res = self.client.get('/api/dashboard/summary/')
        self.assertEqual(dash_res.data['total_expenses'], '60.00')
        self.assertEqual(dash_res.data['recent_transactions'][0]['category'], 'Entertainment')

    def test_non_existent_records_return_404_across_all_endpoints(self):
        """Accessing or mutating non-existent IDs returns 404 across all endpoints"""
        self.auth_user1()
        non_existent_id = 999999

        for endpoint in ['expenses', 'incomes', 'budgets']:
            self.assertEqual(self.client.get(f'/api/{endpoint}/{non_existent_id}/').status_code, status.HTTP_404_NOT_FOUND)
            self.assertEqual(self.client.put(f'/api/{endpoint}/{non_existent_id}/', {'amount': '10.00'}, format='json').status_code, status.HTTP_404_NOT_FOUND)
            self.assertEqual(self.client.patch(f'/api/{endpoint}/{non_existent_id}/', {'amount': '10.00'}, format='json').status_code, status.HTTP_404_NOT_FOUND)
            self.assertEqual(self.client.delete(f'/api/{endpoint}/{non_existent_id}/').status_code, status.HTTP_404_NOT_FOUND)

    def test_unauthenticated_requests_return_401_across_all_milestone2_endpoints(self):
        """Unauthenticated requests are rejected with 401 across all Milestone 2 endpoints"""
        self.clear_auth()

        endpoints = [
            '/api/expenses/',
            '/api/incomes/',
            '/api/budgets/',
            '/api/dashboard/summary/',
            '/api/budgets/summary/',
        ]

        for url in endpoints:
            self.assertEqual(self.client.get(url).status_code, status.HTTP_401_UNAUTHORIZED, f"GET {url} did not return 401")
            self.assertEqual(self.client.post(url, {}, format='json').status_code, status.HTTP_401_UNAUTHORIZED, f"POST {url} did not return 401")

    def test_budget_upsert_logic_for_same_month_and_category(self):
        """
        Setting a budget again for the same category in the same month
        updates the existing allocation seamlessly without creating duplicates.
        """
        self.auth_user1()

        # First allocation
        res1 = self.client.post('/api/budgets/', {
            'category': 'Food',
            'amount': '200.00',
            'month': '2026-09-01'
        }, format='json')
        self.assertEqual(res1.status_code, status.HTTP_201_CREATED)
        self.assertEqual(res1.data['amount'], '200.00')

        # Second allocation for same month & category -> updates existing
        res2 = self.client.post('/api/budgets/', {
            'category': 'Food',
            'amount': '350.00',
            'month': '2026-09-01'
        }, format='json')
        self.assertEqual(res2.status_code, status.HTTP_200_OK)
        self.assertEqual(res2.data['amount'], '350.00')

        # Verify DB has only 1 row for Food in 2026-09
        count = Budget.objects.filter(user=self.user1, category='Food', month__year=2026, month__month=9).count()
        self.assertEqual(count, 1)

    def test_comprehensive_invalid_inputs_and_edge_cases(self):
        """
        Test all edge-case failure inputs:
        - Non-numeric strings
        - Invalid date formats
        - Missing required keys
        """
        self.auth_user1()

        # String amounts
        self.assertEqual(self.client.post('/api/expenses/', {'category': 'Food', 'amount': 'not-a-number', 'date': '2026-09-01'}, format='json').status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(self.client.post('/api/incomes/', {'source': 'Pocket Money', 'amount': 'abc', 'date': '2026-09-01'}, format='json').status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(self.client.post('/api/budgets/', {'category': 'Food', 'amount': 'xyz', 'month': '2026-09-01'}, format='json').status_code, status.HTTP_400_BAD_REQUEST)

        # Invalid date formats
        self.assertEqual(self.client.post('/api/expenses/', {'category': 'Food', 'amount': '20.00', 'date': 'not-a-date'}, format='json').status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(self.client.post('/api/incomes/', {'source': 'Pocket Money', 'amount': '20.00', 'date': '99/99/9999'}, format='json').status_code, status.HTTP_400_BAD_REQUEST)

        # Missing required keys
        self.assertEqual(self.client.post('/api/expenses/', {}, format='json').status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(self.client.post('/api/incomes/', {}, format='json').status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(self.client.post('/api/budgets/', {}, format='json').status_code, status.HTTP_400_BAD_REQUEST)


class Milestone2BoundaryAndDefectTests(APITestCase):
    """
    Boundary Value Analysis and Regression Verification for Milestone 2:
    - Lower boundary amounts (0.00, -0.01 vs 0.01)
    - Upper boundary amounts (99,999,999.99 vs 100,000,000.00)
    - Decimal precision (2 decimal places vs 3 decimal places)
    - Password length boundaries (7 chars vs 8 chars)
    - Description text length boundaries (255 chars vs 256 chars)
    - Date boundaries (leap year, month boundaries)
    """

    def setUp(self):
        self.user = User.objects.create_user(
            username='boundary_user',
            email='boundary@example.com',
            password='Password123!'
        )
        self.token = RefreshToken.for_user(self.user).access_token
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token}')

    def test_boundary_amount_lower_limit(self):
        """
        Below limit: 0.00, -0.01 -> 400 Bad Request
        At limit: 0.01 -> 201 Created
        """
        # Below limit (0.00)
        res_zero = self.client.post('/api/expenses/', {'category': 'Food', 'amount': '0.00', 'date': '2026-09-01'}, format='json')
        self.assertEqual(res_zero.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('amount', res_zero.data)

        # Below limit (-0.01)
        res_neg = self.client.post('/api/expenses/', {'category': 'Food', 'amount': '-0.01', 'date': '2026-09-01'}, format='json')
        self.assertEqual(res_neg.status_code, status.HTTP_400_BAD_REQUEST)

        # At limit (0.01)
        res_min = self.client.post('/api/expenses/', {'category': 'Food', 'amount': '0.01', 'date': '2026-09-01'}, format='json')
        self.assertEqual(res_min.status_code, status.HTTP_201_CREATED)
        self.assertEqual(res_min.data['amount'], '0.01')

    def test_boundary_amount_upper_limit(self):
        """
        At limit: 99999999.99 (10 digits total, 2 decimal places) -> 201 Created
        Above limit: 100000000.00 (11 digits total) -> 400 Bad Request
        """
        # At upper limit
        res_max = self.client.post('/api/expenses/', {
            'category': 'Food',
            'amount': '99999999.99',
            'date': '2026-09-01'
        }, format='json')
        self.assertEqual(res_max.status_code, status.HTTP_201_CREATED)

        # Above upper limit
        res_exceed = self.client.post('/api/expenses/', {
            'category': 'Food',
            'amount': '100000000.00',
            'date': '2026-09-01'
        }, format='json')
        self.assertEqual(res_exceed.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('amount', res_exceed.data)

    def test_boundary_amount_decimal_precision(self):
        """
        At limit: 2 decimal places -> 201 Created
        Above limit: 3 decimal places (10.999) -> 400 Bad Request
        """
        # 2 decimal places
        res_valid = self.client.post('/api/incomes/', {
            'source': 'Pocket Money',
            'amount': '50.25',
            'date': '2026-09-01'
        }, format='json')
        self.assertEqual(res_valid.status_code, status.HTTP_201_CREATED)

        # 3 decimal places
        res_invalid = self.client.post('/api/incomes/', {
            'source': 'Pocket Money',
            'amount': '50.255',
            'date': '2026-09-01'
        }, format='json')
        self.assertEqual(res_invalid.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('amount', res_invalid.data)

    def test_boundary_password_length(self):
        """
        Below limit: 7 chars -> 400 Bad Request
        At limit: 8 chars -> 201 Created
        """
        self.client.credentials()  # Unauthenticated for registration

        # 7 chars
        res_7 = self.client.post('/api/auth/register/', {
            'username': 'user_pass_7',
            'email': 'pass7@example.com',
            'password': 'Pass12!'
        }, format='json')
        self.assertEqual(res_7.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('password', res_7.data)

        # 8 chars
        res_8 = self.client.post('/api/auth/register/', {
            'username': 'user_pass_8',
            'email': 'pass8@example.com',
            'password': 'Pass123!'
        }, format='json')
        self.assertEqual(res_8.status_code, status.HTTP_201_CREATED)

    def test_boundary_description_length(self):
        """
        At limit: 255 chars -> 201 Created
        Above limit: 256 chars -> 400 Bad Request
        """
        # 255 chars
        desc_255 = 'A' * 255
        res_255 = self.client.post('/api/expenses/', {
            'category': 'Shopping',
            'amount': '15.00',
            'date': '2026-09-01',
            'description': desc_255
        }, format='json')
        self.assertEqual(res_255.status_code, status.HTTP_201_CREATED)

        # 256 chars
        desc_256 = 'A' * 256
        res_256 = self.client.post('/api/expenses/', {
            'category': 'Shopping',
            'amount': '15.00',
            'date': '2026-09-01',
            'description': desc_256
        }, format='json')
        self.assertEqual(res_256.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('description', res_256.data)

    def test_boundary_date_leap_year_and_month_boundaries(self):
        """
        Valid leap day: 2024-02-29 -> 201 Created
        Month start: 2026-09-01 -> 201 Created
        Month end: 2026-09-30 -> 201 Created
        Invalid leap day in non-leap year: 2026-02-29 -> 400 Bad Request
        """
        # Leap year date
        res_leap = self.client.post('/api/incomes/', {
            'source': 'Scholarship',
            'amount': '500.00',
            'date': '2024-02-29'
        }, format='json')
        self.assertEqual(res_leap.status_code, status.HTTP_201_CREATED)

        # Month boundaries
        res_start = self.client.post('/api/expenses/', {'category': 'Food', 'amount': '10.00', 'date': '2026-09-01'}, format='json')
        self.assertEqual(res_start.status_code, status.HTTP_201_CREATED)

        res_end = self.client.post('/api/expenses/', {'category': 'Food', 'amount': '10.00', 'date': '2026-09-30'}, format='json')
        self.assertEqual(res_end.status_code, status.HTTP_201_CREATED)

        # Invalid non-leap year Feb 29
        res_invalid_leap = self.client.post('/api/expenses/', {'category': 'Food', 'amount': '10.00', 'date': '2026-02-29'}, format='json')
        self.assertEqual(res_invalid_leap.status_code, status.HTTP_400_BAD_REQUEST)


class SavingsGoalManagementTests(APITestCase):
    """
    Comprehensive tests for Milestone 3: Savings Goal Management
    - Goal Creation, Retrieval, Full/Partial Update, Progress Deposit & Deletion
    - Progress Calculation: Progress % = (Current Saved / Target Amount) * 100
    - Target Reached / Completion Status Detection (is_completed: True/False)
    - User Ownership & Strict Data Isolation
    - Validation: Missing names, zero/negative targets, negative progress, invalid IDs
    - Unauthenticated request rejection (401)
    """

    def setUp(self):
        self.user_a = User.objects.create_user(
            username='saver_alice',
            email='alice@example.com',
            password='Password123!'
        )
        self.user_b = User.objects.create_user(
            username='saver_bob',
            email='bob@example.com',
            password='Password123!'
        )
        self.token_a = RefreshToken.for_user(self.user_a).access_token
        self.token_b = RefreshToken.for_user(self.user_b).access_token

        # Baseline goal for Alice
        self.goal_a = SavingsGoal.objects.create(
            user=self.user_a,
            name='Emergency Fund',
            target_amount='50000.00',
            current_amount='15000.00',
            target_date=date(2026, 12, 31)
        )

        # Baseline goal for Bob
        self.goal_b = SavingsGoal.objects.create(
            user=self.user_b,
            name='New Laptop',
            target_amount='60000.00',
            current_amount='20000.00',
            target_date=date(2026, 11, 30)
        )

    def auth_alice(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token_a}')

    def auth_bob(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token_b}')

    def clear_auth(self):
        self.client.credentials()

    # --- 1. Creation & Progress Calculation ---
    def test_create_savings_goal_success(self):
        """Authenticated user creates a savings goal with target and initial amount"""
        self.auth_alice()
        data = {
            'name': 'Trip to Japan',
            'target_amount': '100000.00',
            'current_amount': '25000.00',
            'target_date': '2027-05-01'
        }
        res = self.client.post('/api/savings-goals/', data, format='json')

        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertEqual(res.data['name'], 'Trip to Japan')
        self.assertEqual(res.data['target_amount'], '100000.00')
        self.assertEqual(res.data['current_amount'], '25000.00')
        self.assertEqual(res.data['progress_percentage'], 25.0)
        self.assertEqual(res.data['remaining_amount'], '75000.00')
        self.assertFalse(res.data['is_completed'])

        # Verify DB association
        db_goal = SavingsGoal.objects.get(id=res.data['id'])
        self.assertEqual(db_goal.user, self.user_a)

    def test_savings_goal_progress_calculation_standard_example(self):
        """
        Verify standard requirement calculation:
        Target Amount = ₹50,000
        Saved Amount  = ₹15,000
        Progress      = 30%
        """
        self.auth_alice()
        res = self.client.get(f'/api/savings-goals/{self.goal_a.id}/')

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data['name'], 'Emergency Fund')
        self.assertEqual(res.data['target_amount'], '50000.00')
        self.assertEqual(res.data['current_amount'], '15000.00')
        self.assertEqual(res.data['progress_percentage'], 30.0)
        self.assertEqual(res.data['remaining_amount'], '35000.00')
        self.assertFalse(res.data['is_completed'])

    def test_savings_goal_completion_state_reached_and_exceeded(self):
        """
        Target reached (saved >= target) flags is_completed = True, remaining = 0.00.
        When target is exceeded, progress > 100% and is_completed = True.
        """
        self.auth_alice()

        # Update saved amount to exactly reach target (₹50,000)
        patch_res = self.client.patch(f'/api/savings-goals/{self.goal_a.id}/', {
            'current_amount': '50000.00'
        }, format='json')
        self.assertEqual(patch_res.status_code, status.HTTP_200_OK)
        self.assertEqual(patch_res.data['progress_percentage'], 100.0)
        self.assertEqual(patch_res.data['remaining_amount'], '0.00')
        self.assertTrue(patch_res.data['is_completed'])

        # Exceed target (₹55,000 saved out of ₹50,000)
        patch_res2 = self.client.patch(f'/api/savings-goals/{self.goal_a.id}/', {
            'current_amount': '55000.00'
        }, format='json')
        self.assertEqual(patch_res2.status_code, status.HTTP_200_OK)
        self.assertEqual(patch_res2.data['progress_percentage'], 110.0)
        self.assertEqual(patch_res2.data['remaining_amount'], '0.00')
        self.assertTrue(patch_res2.data['is_completed'])

    def test_savings_goal_progress_deposit_via_patch(self):
        """Updating progress amount via PATCH accurately reflects new progress percentage"""
        self.auth_alice()

        # Alice starts with 15,000 / 50,000 (30%). Deposits 10,000 -> 25,000 / 50,000 (50%).
        res = self.client.patch(f'/api/savings-goals/{self.goal_a.id}/', {
            'current_amount': '25000.00'
        }, format='json')

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data['current_amount'], '25000.00')
        self.assertEqual(res.data['progress_percentage'], 50.0)
        self.assertEqual(res.data['remaining_amount'], '25000.00')
        self.assertFalse(res.data['is_completed'])

    # --- 2. Full Update, Listing & Detail ---
    def test_savings_goal_full_update_via_put(self):
        """Full update via PUT replaces goal attributes and recalculates progress"""
        self.auth_alice()
        data = {
            'name': 'Updated Emergency Reserve',
            'target_amount': '80000.00',
            'current_amount': '40000.00',
            'target_date': '2027-01-01'
        }
        res = self.client.put(f'/api/savings-goals/{self.goal_a.id}/', data, format='json')

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data['name'], 'Updated Emergency Reserve')
        self.assertEqual(res.data['target_amount'], '80000.00')
        self.assertEqual(res.data['current_amount'], '40000.00')
        self.assertEqual(res.data['progress_percentage'], 50.0)
        self.assertEqual(res.data['target_date'], '2027-01-01')

    def test_savings_goal_list_retrieval(self):
        """User retrieves all their own savings goals ordered by id descending"""
        self.auth_alice()

        # Add second goal for Alice
        SavingsGoal.objects.create(
            user=self.user_a,
            name='Certification Course',
            target_amount='10000.00',
            current_amount='5000.00'
        )

        res = self.client.get('/api/savings-goals/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(len(res.data), 2)
        goal_names = [g['name'] for g in res.data]
        self.assertIn('Certification Course', goal_names)
        self.assertIn('Emergency Fund', goal_names)

    # --- 3. User Ownership & Data Isolation ---
    def test_savings_goal_user_data_isolation(self):
        """User Alice cannot view, update, or delete User Bob's savings goal"""
        self.auth_alice()

        # Alice tries to view Bob's goal -> 404
        res_get = self.client.get(f'/api/savings-goals/{self.goal_b.id}/')
        self.assertEqual(res_get.status_code, status.HTTP_404_NOT_FOUND)

        # Alice tries to update Bob's goal -> 404
        res_put = self.client.put(f'/api/savings-goals/{self.goal_b.id}/', {
            'name': 'Hacked Goal',
            'target_amount': '1.00',
            'current_amount': '1.00'
        }, format='json')
        self.assertEqual(res_put.status_code, status.HTTP_404_NOT_FOUND)

        # Alice tries to patch Bob's goal -> 404
        res_patch = self.client.patch(f'/api/savings-goals/{self.goal_b.id}/', {
            'current_amount': '0.00'
        }, format='json')
        self.assertEqual(res_patch.status_code, status.HTTP_404_NOT_FOUND)

        # Alice tries to delete Bob's goal -> 404
        res_del = self.client.delete(f'/api/savings-goals/{self.goal_b.id}/')
        self.assertEqual(res_del.status_code, status.HTTP_404_NOT_FOUND)

        # Verify Bob's goal is intact in DB
        self.goal_b.refresh_from_db()
        self.assertEqual(self.goal_b.name, 'New Laptop')
        self.assertEqual(str(self.goal_b.current_amount), '20000.00')

    # --- 4. Deletion ---
    def test_savings_goal_deletion_success(self):
        """User can delete their own savings goal"""
        self.auth_alice()
        res = self.client.delete(f'/api/savings-goals/{self.goal_a.id}/')

        self.assertEqual(res.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(SavingsGoal.objects.filter(id=self.goal_a.id).exists())

    # --- 5. Validation & Boundary Handling ---
    def test_savings_goal_validation_errors(self):
        """
        Rejects invalid inputs with HTTP 400:
        - Missing or blank name
        - Zero or negative target amounts
        - Negative progress / saved amounts
        """
        self.auth_alice()

        # 1. Missing name
        res_no_name = self.client.post('/api/savings-goals/', {
            'target_amount': '5000.00',
            'current_amount': '1000.00'
        }, format='json')
        self.assertEqual(res_no_name.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('name', res_no_name.data)

        # 2. Blank name
        res_blank_name = self.client.post('/api/savings-goals/', {
            'name': '   ',
            'target_amount': '5000.00'
        }, format='json')
        self.assertEqual(res_blank_name.status_code, status.HTTP_400_BAD_REQUEST)

        # 3. Zero target amount
        res_zero_target = self.client.post('/api/savings-goals/', {
            'name': 'Zero Target Goal',
            'target_amount': '0.00'
        }, format='json')
        self.assertEqual(res_zero_target.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('target_amount', res_zero_target.data)

        # 4. Negative target amount
        res_neg_target = self.client.post('/api/savings-goals/', {
            'name': 'Negative Target Goal',
            'target_amount': '-500.00'
        }, format='json')
        self.assertEqual(res_neg_target.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('target_amount', res_neg_target.data)

        # 5. Negative current amount
        res_neg_saved = self.client.post('/api/savings-goals/', {
            'name': 'Negative Saved Goal',
            'target_amount': '5000.00',
            'current_amount': '-100.00'
        }, format='json')
        self.assertEqual(res_neg_saved.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('current_amount', res_neg_saved.data)

    def test_savings_goal_non_existent_record_returns_404(self):
        """Mutating or viewing a non-existent goal ID returns 404"""
        self.auth_alice()
        non_existent_id = 99999

        self.assertEqual(self.client.get(f'/api/savings-goals/{non_existent_id}/').status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(self.client.put(f'/api/savings-goals/{non_existent_id}/', {'name': 'Test', 'target_amount': '100.00'}, format='json').status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(self.client.patch(f'/api/savings-goals/{non_existent_id}/', {'current_amount': '50.00'}, format='json').status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(self.client.delete(f'/api/savings-goals/{non_existent_id}/').status_code, status.HTTP_404_NOT_FOUND)

    # --- 6. Unauthenticated Requests ---
    def test_savings_goal_unauthenticated_requests_return_401(self):
        """Unauthenticated requests to savings goals endpoints return 401 Unauthorized"""
        self.clear_auth()

        self.assertEqual(self.client.get('/api/savings-goals/').status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(self.client.post('/api/savings-goals/', {'name': 'Test', 'target_amount': '100.00'}, format='json').status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(self.client.get(f'/api/savings-goals/{self.goal_a.id}/').status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(self.client.put(f'/api/savings-goals/{self.goal_a.id}/', {'name': 'Test', 'target_amount': '100.00'}, format='json').status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(self.client.patch(f'/api/savings-goals/{self.goal_a.id}/', {'current_amount': '50.00'}, format='json').status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(self.client.delete(f'/api/savings-goals/{self.goal_a.id}/').status_code, status.HTTP_401_UNAUTHORIZED)


class NotificationSystemTests(APITestCase):
    """
    Test suite for Milestone 4: Notification System
    - Model creation, defaults, ordering
    - API retrieval with unread filters and unread count
    - Mark as read (single and all)
    - User ownership and tenant isolation
    - Unauthenticated request rejection (401)
    - Budget limit alert triggers & threshold checks
    - Savings goal milestone triggers (50% & 100%)
    - Duplicate alert prevention
    """
    def setUp(self):
        # Create user Alice
        self.user_alice = User.objects.create_user(
            username='alice_notif',
            email='alice_notif@example.com',
            password='Password123!'
        )
        refresh_alice = RefreshToken.for_user(self.user_alice)
        self.token_alice = str(refresh_alice.access_token)

        # Create user Bob
        self.user_bob = User.objects.create_user(
            username='bob_notif',
            email='bob_notif@example.com',
            password='Password123!'
        )
        refresh_bob = RefreshToken.for_user(self.user_bob)
        self.token_bob = str(refresh_bob.access_token)

        # Seed notification for Alice
        self.notif_alice_1 = Notification.objects.create(
            user=self.user_alice,
            notification_type='budget_alert',
            title='Budget Alert: Food Exceeded',
            message='You have exceeded your Food budget.',
            related_context='seed:alice_budget_alert',
            is_read=False
        )
        self.notif_alice_2 = Notification.objects.create(
            user=self.user_alice,
            notification_type='goal_milestone',
            title='Savings Milestone: Laptop (50%)',
            message='You reached 50% of your Laptop savings goal.',
            related_context='seed:alice_goal_milestone',
            is_read=True
        )

        # Seed notification for Bob
        self.notif_bob_1 = Notification.objects.create(
            user=self.user_bob,
            notification_type='budget_alert',
            title='Budget Alert: Travel Exceeded',
            message='You have exceeded your Travel budget.',
            related_context='seed:bob_budget_alert',
            is_read=False
        )

    def auth_alice(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token_alice}')

    def auth_bob(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token_bob}')

    def clear_auth(self):
        self.client.credentials()

    # --- 1. Model Structure & Defaults ---
    def test_notification_model_creation_and_defaults(self):
        """Notification model structure, fields, and defaults"""
        notif = Notification.objects.create(
            user=self.user_alice,
            title='Test Notification',
            message='This is a test notification message.'
        )
        self.assertEqual(notif.user, self.user_alice)
        self.assertEqual(notif.notification_type, 'system')
        self.assertFalse(notif.is_read)
        self.assertIsNotNone(notif.created_at)
        self.assertIn('alice_notif', str(notif))

    # --- 2. API Retrieval & User Isolation ---
    def test_get_notifications_user_isolation(self):
        """User only retrieves their own notifications"""
        self.auth_alice()
        response = self.client.get('/api/notifications/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        
        # Alice should see 2 notifications
        self.assertEqual(len(response.data), 2)
        notif_ids = [n['id'] for n in response.data]
        self.assertIn(self.notif_alice_1.id, notif_ids)
        self.assertIn(self.notif_alice_2.id, notif_ids)
        self.assertNotIn(self.notif_bob_1.id, notif_ids)

    def test_get_unread_notifications_filter(self):
        """Filter notifications by unread=true"""
        self.auth_alice()
        response = self.client.get('/api/notifications/?unread=true')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        
        # Only unread notification (notif_alice_1) should be returned
        self.assertEqual(len(response.data), 1)
        self.assertEqual(response.data[0]['id'], self.notif_alice_1.id)
        self.assertFalse(response.data[0]['is_read'])

    def test_get_unread_count(self):
        """GET /api/notifications/unread-count/ returns accurate count"""
        self.auth_alice()
        response = self.client.get('/api/notifications/unread-count/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['unread_count'], 1)

        self.auth_bob()
        response_bob = self.client.get('/api/notifications/unread-count/')
        self.assertEqual(response_bob.status_code, status.HTTP_200_OK)
        self.assertEqual(response_bob.data['unread_count'], 1)

    # --- 3. Mark as Read & User Isolation ---
    def test_mark_single_notification_as_read(self):
        """Mark a single notification as read via PATCH /read/ and PATCH /"""
        self.auth_alice()
        
        # Test PATCH /api/notifications/<id>/read/
        response = self.client.patch(f'/api/notifications/{self.notif_alice_1.id}/read/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.data['is_read'])

        self.notif_alice_1.refresh_from_db()
        self.assertTrue(self.notif_alice_1.is_read)

        # Unread count should now be 0
        count_res = self.client.get('/api/notifications/unread-count/')
        self.assertEqual(count_res.data['unread_count'], 0)

    def test_user_cannot_access_or_modify_other_users_notification(self):
        """Alice cannot view, modify, mark as read, or delete Bob's notification"""
        self.auth_alice()

        # View Bob's notification
        self.assertEqual(self.client.get(f'/api/notifications/{self.notif_bob_1.id}/').status_code, status.HTTP_404_NOT_FOUND)
        # Mark Bob's notification as read
        self.assertEqual(self.client.patch(f'/api/notifications/{self.notif_bob_1.id}/read/').status_code, status.HTTP_404_NOT_FOUND)
        # Patch Bob's notification
        self.assertEqual(self.client.patch(f'/api/notifications/{self.notif_bob_1.id}/', {'is_read': True}, format='json').status_code, status.HTTP_404_NOT_FOUND)
        # Delete Bob's notification
        self.assertEqual(self.client.delete(f'/api/notifications/{self.notif_bob_1.id}/').status_code, status.HTTP_404_NOT_FOUND)

        # Bob's notification remains untouched
        self.notif_bob_1.refresh_from_db()
        self.assertFalse(self.notif_bob_1.is_read)

    def test_mark_all_notifications_as_read(self):
        """POST /api/notifications/mark-all-read/ marks all user's notifications as read"""
        self.auth_alice()

        # Add another unread notification for Alice
        Notification.objects.create(
            user=self.user_alice,
            title='Another Alert',
            message='Another message',
            is_read=False
        )

        self.assertEqual(Notification.objects.filter(user=self.user_alice, is_read=False).count(), 2)

        response = self.client.post('/api/notifications/mark-all-read/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['updated_count'], 2)

        # Alice has 0 unread
        self.assertEqual(Notification.objects.filter(user=self.user_alice, is_read=False).count(), 0)
        # Bob's notification is untouched
        self.assertEqual(Notification.objects.filter(user=self.user_bob, is_read=False).count(), 1)

    def test_delete_notification(self):
        """DELETE /api/notifications/<id>/ deletes user's notification"""
        self.auth_alice()
        response = self.client.delete(f'/api/notifications/{self.notif_alice_1.id}/')
        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(Notification.objects.filter(id=self.notif_alice_1.id).exists())

    def test_unauthenticated_requests_return_401(self):
        """Unauthenticated requests to notification endpoints return 401"""
        self.clear_auth()
        self.assertEqual(self.client.get('/api/notifications/').status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(self.client.get('/api/notifications/unread-count/').status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(self.client.post('/api/notifications/mark-all-read/').status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(self.client.get(f'/api/notifications/{self.notif_alice_1.id}/').status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(self.client.patch(f'/api/notifications/{self.notif_alice_1.id}/read/').status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(self.client.delete(f'/api/notifications/{self.notif_alice_1.id}/').status_code, status.HTTP_401_UNAUTHORIZED)

    # --- 4. Budget Alert Triggers & Thresholds ---
    def test_budget_limit_alert_triggered_when_spending_exceeds_budget(self):
        """Creating an expense exceeding budget triggers budget_alert notification"""
        self.auth_alice()
        target_month = date.today().replace(day=1)

        # Set Food budget of ₹1000
        Budget.objects.create(
            user=self.user_alice,
            category='Food',
            amount=Decimal('1000.00'),
            month=target_month
        )

        initial_notif_count = Notification.objects.filter(user=self.user_alice, notification_type='budget_alert').count()

        # 1. Expense below budget (₹400) -> No notification
        res1 = self.client.post('/api/expenses/', {
            'category': 'Food',
            'amount': '400.00',
            'date': str(target_month),
            'description': 'Groceries'
        }, format='json')
        self.assertEqual(res1.status_code, status.HTTP_201_CREATED)
        self.assertEqual(Notification.objects.filter(user=self.user_alice, notification_type='budget_alert').count(), initial_notif_count)

        # 2. Expense that crosses budget (₹700 -> Total ₹1100 >= ₹1000) -> Triggers notification!
        res2 = self.client.post('/api/expenses/', {
            'category': 'Food',
            'amount': '700.00',
            'date': str(target_month),
            'description': 'Restaurant dinner'
        }, format='json')
        self.assertEqual(res2.status_code, status.HTTP_201_CREATED)

        latest_notif = Notification.objects.filter(user=self.user_alice, notification_type='budget_alert').order_by('-id').first()
        self.assertIsNotNone(latest_notif)
        self.assertIn('Food Exceeded', latest_notif.title)
        self.assertIn('1100.00', latest_notif.message)
        self.assertIn('1000.00', latest_notif.message)
        self.assertFalse(latest_notif.is_read)

    def test_budget_alert_deduplication(self):
        """Subsequent expenses in the same month do not create duplicate alert notifications"""
        self.auth_alice()
        target_month = date.today().replace(day=1)

        budget = Budget.objects.create(
            user=self.user_alice,
            category='Shopping',
            amount=Decimal('500.00'),
            month=target_month
        )

        expected_context = f"budget:{budget.id}:{target_month.year}-{target_month.month:02d}"

        # First expense exceeds budget
        self.client.post('/api/expenses/', {
            'category': 'Shopping',
            'amount': '600.00',
            'date': str(target_month),
            'description': 'Clothes'
        }, format='json')

        alert_count_after_first = Notification.objects.filter(
            user=self.user_alice,
            notification_type='budget_alert',
            related_context=expected_context
        ).count()
        self.assertEqual(alert_count_after_first, 1)

        # Second expense in same category and month
        self.client.post('/api/expenses/', {
            'category': 'Shopping',
            'amount': '200.00',
            'date': str(target_month),
            'description': 'Shoes'
        }, format='json')

        alert_count_after_second = Notification.objects.filter(
            user=self.user_alice,
            notification_type='budget_alert',
            related_context=expected_context
        ).count()

        # Should not duplicate
        self.assertEqual(alert_count_after_second, 1)

    def test_budget_alert_at_limit_and_isolated_by_category_and_user(self):
        """An exact limit triggers once and ignores other categories and users."""
        self.auth_alice()
        target_month = date.today().replace(day=1)

        food_budget = Budget.objects.create(
            user=self.user_alice,
            category='Food',
            amount=Decimal('1000.00'),
            month=target_month
        )
        Budget.objects.create(
            user=self.user_alice,
            category='Travel',
            amount=Decimal('1000.00'),
            month=target_month
        )
        Budget.objects.create(
            user=self.user_bob,
            category='Food',
            amount=Decimal('1000.00'),
            month=target_month
        )

        response = self.client.post('/api/expenses/', {
            'category': 'Food',
            'amount': '1000.00',
            'date': str(target_month),
            'description': 'Exact budget limit'
        }, format='json')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)

        context = f"budget:{food_budget.id}:{target_month.year}-{target_month.month:02d}"
        self.assertEqual(Notification.objects.filter(
            user=self.user_alice,
            notification_type='budget_alert',
            related_context=context
        ).count(), 1)
        self.assertFalse(Notification.objects.filter(
            user=self.user_alice,
            related_context__contains='warning'
        ).exists())

        self.client.post('/api/expenses/', {
            'category': 'Travel',
            'amount': '900.00',
            'date': str(target_month),
            'description': 'Travel remains below limit'
        }, format='json')
        self.assertEqual(Notification.objects.filter(
            user=self.user_alice,
            notification_type='budget_alert',
            related_context=context
        ).count(), 1)
        self.assertTrue(Notification.objects.filter(
            user=self.user_alice,
            related_context=f"budget:{food_budget.id}:{target_month.year}-{target_month.month:02d}"
        ).exists())

        self.auth_bob()
        self.client.post('/api/expenses/', {
            'category': 'Food',
            'amount': '500.00',
            'date': str(target_month),
            'description': 'Bob has separate spending'
        }, format='json')
        self.assertEqual(Notification.objects.filter(
            user=self.user_bob,
            notification_type='budget_alert'
        ).count(), 1)

    # --- 5. Savings Goal Milestone Triggers ---
    def test_savings_goal_milestone_and_completion_triggers(self):
        """Savings goal triggers 50% milestone alert and 100% completion alert"""
        self.auth_alice()

        # 1. Create goal with ₹5000 target and ₹1000 saved (20%) -> No alert
        res_create = self.client.post('/api/savings-goals/', {
            'name': 'New Tablet',
            'target_amount': '5000.00',
            'current_amount': '1000.00'
        }, format='json')
        self.assertEqual(res_create.status_code, status.HTTP_201_CREATED)
        goal_id = res_create.data['id']

        self.assertFalse(Notification.objects.filter(user=self.user_alice, related_context=f'goal:{goal_id}:50pct').exists())
        self.assertFalse(Notification.objects.filter(user=self.user_alice, related_context=f'goal:{goal_id}:completed').exists())

        # 2. Update to ₹2500 (50% milestone) -> Triggers 50% milestone alert
        res_patch_50 = self.client.patch(f'/api/savings-goals/{goal_id}/', {
            'current_amount': '2500.00'
        }, format='json')
        self.assertEqual(res_patch_50.status_code, status.HTTP_200_OK)

        milestone_50 = Notification.objects.filter(user=self.user_alice, related_context=f'goal:{goal_id}:50pct').first()
        self.assertIsNotNone(milestone_50)
        self.assertIn('50%', milestone_50.title)
        self.assertIn('2500.00', milestone_50.message)

        # 3. Update to ₹5000 (100% completion) -> Triggers goal achieved alert
        res_patch_100 = self.client.patch(f'/api/savings-goals/{goal_id}/', {
            'current_amount': '5000.00'
        }, format='json')
        self.assertEqual(res_patch_100.status_code, status.HTTP_200_OK)

        completed_alert = Notification.objects.filter(user=self.user_alice, related_context=f'goal:{goal_id}:completed').first()
        self.assertIsNotNone(completed_alert)
        self.assertIn('Achieved', completed_alert.title)
        self.assertIn('5000.00', completed_alert.message)

    def test_savings_goal_milestone_deduplication(self):
        """Subsequent updates do not duplicate the same milestone alert"""
        self.auth_alice()

        goal = SavingsGoal.objects.create(
            user=self.user_alice,
            name='Headphones',
            target_amount=Decimal('2000.00'),
            current_amount=Decimal('1000.00') # 50%
        )

        # Initial trigger
        from .notifications import check_and_trigger_goal_alert
        check_and_trigger_goal_alert(goal)

        count_1 = Notification.objects.filter(user=self.user_alice, related_context=f'goal:{goal.id}:50pct').count()
        self.assertEqual(count_1, 1)

        # Repeated trigger call
        check_and_trigger_goal_alert(goal)
        count_2 = Notification.objects.filter(user=self.user_alice, related_context=f'goal:{goal.id}:50pct').count()
        self.assertEqual(count_2, 1)

    def test_budget_near_limit_warning_threshold(self):
        """Spending reaching >= 80% and < 100% triggers budget warning alert"""
        self.auth_alice()
        target_month = date.today().replace(day=1)

        budget = Budget.objects.create(
            user=self.user_alice,
            category='Education',
            amount=Decimal('1000.00'),
            month=target_month
        )

        # 1. Below 80% (₹500 = 50%) -> No warning
        res1 = self.client.post('/api/expenses/', {
            'category': 'Education',
            'amount': '500.00',
            'date': str(target_month),
            'description': 'Books'
        }, format='json')
        self.assertEqual(res1.status_code, status.HTTP_201_CREATED)
        self.assertFalse(Notification.objects.filter(user=self.user_alice, related_context__contains='warning').exists())

        # 2. Crosses 80% (₹350 -> Total ₹850 = 85%) -> Triggers Near Limit warning!
        res2 = self.client.post('/api/expenses/', {
            'category': 'Education',
            'amount': '350.00',
            'date': str(target_month),
            'description': 'Course fee'
        }, format='json')
        self.assertEqual(res2.status_code, status.HTTP_201_CREATED)

        warn_notif = Notification.objects.filter(user=self.user_alice, related_context=f"budget:{budget.id}:{target_month.year}-{target_month.month:02d}:warning").first()
        self.assertIsNotNone(warn_notif)
        self.assertIn('Near Limit', warn_notif.title)
        self.assertIn('85.0%', warn_notif.message)


class AnalyticsAndReportingTests(APITestCase):
    """
    Test suite for Analytics APIs & Monthly Financial Reports:
    - Overview metrics calculation with known records (income, expenses, net savings, savings rate)
    - Category-wise expense distribution & percentage math
    - Date, month, and year filtering
    - Zero values & empty data safety (no DivisionByZero)
    - Monthly trends generation (6/12 months)
    - Report generation workflow & database record matching
    - Report CSV export format & download
    - Report JSON export
    - Notification trigger upon report generation
    - Tenant isolation: User A cannot view/export/delete User B's analytics or reports
    - Unauthenticated request rejection (401)
    """
    def setUp(self):
        self.alice = User.objects.create_user(
            username='alice_analytics',
            email='alice_analytics@example.com',
            password='Password123!'
        )
        refresh_alice = RefreshToken.for_user(self.alice)
        self.token_alice = str(refresh_alice.access_token)

        self.bob = User.objects.create_user(
            username='bob_analytics',
            email='bob_analytics@example.com',
            password='Password123!'
        )
        refresh_bob = RefreshToken.for_user(self.bob)
        self.token_bob = str(refresh_bob.access_token)

        # Alice's seeded data for September 2026
        # Incomes: ₹5000 (Pocket Money) + ₹15000 (Salary) = ₹20000
        Income.objects.create(user=self.alice, source='Pocket Money', amount=Decimal('5000.00'), date=date(2026, 9, 5))
        Income.objects.create(user=self.alice, source='Salary', amount=Decimal('15000.00'), date=date(2026, 9, 25))

        # Expenses: ₹3000 (Food) + ₹2000 (Travel) = ₹5000
        Expense.objects.create(user=self.alice, category='Food', amount=Decimal('3000.00'), date=date(2026, 9, 10))
        Expense.objects.create(user=self.alice, category='Travel', amount=Decimal('2000.00'), date=date(2026, 9, 15))

        # Budgets for September 2026
        Budget.objects.create(user=self.alice, category='Food', amount=Decimal('4000.00'), month=date(2026, 9, 1))
        Budget.objects.create(user=self.alice, category='Travel', amount=Decimal('2500.00'), month=date(2026, 9, 1))

        # Savings Goal
        SavingsGoal.objects.create(
            user=self.alice,
            name='Emergency Fund',
            target_amount=Decimal('50000.00'),
            current_amount=Decimal('10000.00')
        )

        # Bob's seeded data
        Income.objects.create(user=self.bob, source='Salary', amount=Decimal('30000.00'), date=date(2026, 9, 1))
        Expense.objects.create(user=self.bob, category='Entertainment', amount=Decimal('8000.00'), date=date(2026, 9, 2))

    def auth_alice(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token_alice}')

    def auth_bob(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token_bob}')

    def clear_auth(self):
        self.client.credentials()

    # --- 1. Analytics Overview & Calculations ---
    def test_analytics_overview_known_calculations(self):
        """Analytics overview accurately calculates totals, net savings, and savings rate"""
        self.auth_alice()
        response = self.client.get('/api/analytics/overview/?month=2026-09')
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        data = response.data
        self.assertEqual(data['total_income'], '20000.00')
        self.assertEqual(data['total_expenses'], '5000.00')
        self.assertEqual(data['net_savings'], '15000.00')
        # Savings rate = (15000 / 20000) * 100 = 75.0%
        self.assertEqual(data['savings_rate'], 75.0)
        self.assertEqual(data['income_count'], 2)
        self.assertEqual(data['expense_count'], 2)

        # Category breakdown: Food (3000 = 60%), Travel (2000 = 40%)
        cats = {c['category']: c for c in data['categories']}
        self.assertEqual(cats['Food']['total_amount'], '3000.00')
        self.assertEqual(cats['Food']['percentage'], 60.0)
        self.assertEqual(cats['Travel']['total_amount'], '2000.00')
        self.assertEqual(cats['Travel']['percentage'], 40.0)

        # Budget performance: Total budgeted = 4000 + 2500 = 6500, spent = 5000, utilization = (5000/6500)*100 = 76.92%
        self.assertEqual(data['budget']['total_budgeted'], '6500.00')
        self.assertEqual(data['budget']['total_spent'], '5000.00')
        self.assertEqual(data['budget']['utilization_percentage'], 76.92)

    def test_analytics_refreshes_after_income_and_expense_transactions(self):
        """Analytics re-queries reflect newly saved transactions and match the database."""
        self.auth_alice()

        income_response = self.client.post('/api/incomes/', {
            'source': 'Other',
            'description': 'Integration test income',
            'amount': '1250.00',
            'date': '2026-09-20',
        }, format='json')
        expense_response = self.client.post('/api/expenses/', {
            'category': 'Food',
            'description': 'Integration test expense',
            'amount': '250.00',
            'date': '2026-09-21',
        }, format='json')
        self.assertEqual(income_response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(expense_response.status_code, status.HTTP_201_CREATED)

        overview_response = self.client.get('/api/analytics/overview/?month=2026-09')
        self.assertEqual(overview_response.status_code, status.HTTP_200_OK)
        overview = overview_response.data

        db_income = sum(record.amount for record in Income.objects.filter(user=self.alice, date__year=2026, date__month=9))
        db_expenses = sum(record.amount for record in Expense.objects.filter(user=self.alice, date__year=2026, date__month=9))
        self.assertEqual(overview['total_income'], f'{db_income:.2f}')
        self.assertEqual(overview['total_expenses'], f'{db_expenses:.2f}')
        categories = {item['category']: item for item in overview['categories']}
        self.assertEqual(categories['Food']['total_amount'], '3250.00')

        trends_response = self.client.get('/api/analytics/trends/?year=2026')
        self.assertEqual(trends_response.status_code, status.HTTP_200_OK)
        september = next(item for item in trends_response.data['trends'] if item['month'] == '2026-09')
        self.assertEqual(september['total_income'], f'{db_income:.2f}')
        self.assertEqual(september['total_expenses'], f'{db_expenses:.2f}')

    def test_analytics_refreshes_after_expense_changes_and_savings_progress(self):
        """Overview and trends reflect expense edits/deletes and updated savings goals."""
        self.auth_alice()

        food_expense = Expense.objects.filter(user=self.alice, category='Food').first()
        update_response = self.client.patch(f'/api/expenses/{food_expense.id}/', {
            'amount': '3500.00',
        }, format='json')
        self.assertEqual(update_response.status_code, status.HTTP_200_OK)

        overview_response = self.client.get('/api/analytics/overview/?month=2026-09')
        self.assertEqual(overview_response.data['total_expenses'], '5500.00')
        food_category = next(item for item in overview_response.data['categories'] if item['category'] == 'Food')
        self.assertEqual(food_category['raw_total'], 3500.0)

        delete_response = self.client.delete(f"/api/expenses/{update_response.data['id']}/")
        self.assertEqual(delete_response.status_code, status.HTTP_204_NO_CONTENT)

        goal = SavingsGoal.objects.get(user=self.alice, name='Emergency Fund')
        savings_response = self.client.patch(f'/api/savings-goals/{goal.id}/', {
            'current_amount': '25000.00',
        }, format='json')
        self.assertEqual(savings_response.status_code, status.HTTP_200_OK)

        overview_response = self.client.get('/api/analytics/overview/?month=2026-09')
        self.assertEqual(overview_response.data['total_expenses'], '2000.00')
        self.assertEqual(overview_response.data['savings_goals']['total_saved_amount'], '25000.00')
        self.assertEqual(overview_response.data['savings_goals']['overall_progress_percentage'], 50.0)

        trends_response = self.client.get('/api/analytics/trends/?year=2026')
        september = next(item for item in trends_response.data['trends'] if item['month'] == '2026-09')
        self.assertEqual(september['total_expenses'], '2000.00')

    def test_milestone_three_full_financial_workflow(self):
        """Budget alerts, goal milestones, analytics, reports, and ownership work together."""
        self.auth_alice()
        target_date = '2026-09-22'

        income_response = self.client.post('/api/incomes/', {
            'source': 'Other',
            'description': 'Workflow income',
            'amount': '1200.00',
            'date': target_date,
        }, format='json')
        self.assertEqual(income_response.status_code, status.HTTP_201_CREATED)

        warning_expense = self.client.post('/api/expenses/', {
            'category': 'Food',
            'description': 'Cross warning threshold',
            'amount': '200.00',
            'date': target_date,
        }, format='json')
        self.assertEqual(warning_expense.status_code, status.HTTP_201_CREATED)
        budget = Budget.objects.get(user=self.alice, category='Food', month=date(2026, 9, 1))
        warning_context = f'budget:{budget.id}:2026-09:warning'
        self.assertEqual(Notification.objects.filter(
            user=self.alice,
            related_context=warning_context,
            is_read=False,
        ).count(), 1)

        limit_expense = self.client.post('/api/expenses/', {
            'category': 'Food',
            'description': 'Reach budget limit',
            'amount': '800.00',
            'date': target_date,
        }, format='json')
        self.assertEqual(limit_expense.status_code, status.HTTP_201_CREATED)
        alert_context = f'budget:{budget.id}:2026-09'
        self.assertEqual(Notification.objects.filter(
            user=self.alice,
            related_context=alert_context,
            is_read=False,
        ).count(), 1)

        goal_response = self.client.post('/api/savings-goals/', {
            'name': 'Workflow Goal',
            'target_amount': '1000.00',
            'current_amount': '100.00',
        }, format='json')
        self.assertEqual(goal_response.status_code, status.HTTP_201_CREATED)
        goal_id = goal_response.data['id']

        milestone_response = self.client.patch(f'/api/savings-goals/{goal_id}/', {
            'current_amount': '500.00',
        }, format='json')
        self.assertEqual(milestone_response.status_code, status.HTTP_200_OK)
        self.assertEqual(milestone_response.data['progress_percentage'], 50.0)
        self.assertFalse(milestone_response.data['is_completed'])
        self.assertEqual(Notification.objects.filter(
            user=self.alice,
            related_context=f'goal:{goal_id}:50pct',
        ).count(), 1)

        completion_response = self.client.patch(f'/api/savings-goals/{goal_id}/', {
            'current_amount': '1000.00',
        }, format='json')
        self.assertEqual(completion_response.status_code, status.HTTP_200_OK)
        self.assertEqual(completion_response.data['progress_percentage'], 100.0)
        self.assertTrue(completion_response.data['is_completed'])
        self.assertEqual(Notification.objects.filter(
            user=self.alice,
            related_context=f'goal:{goal_id}:completed',
        ).count(), 1)

        edited_response = self.client.patch(f"/api/expenses/{warning_expense.data['id']}/", {
            'amount': '250.00',
        }, format='json')
        self.assertEqual(edited_response.status_code, status.HTTP_200_OK)

        overview_response = self.client.get('/api/analytics/overview/?month=2026-09')
        self.assertEqual(overview_response.status_code, status.HTTP_200_OK)
        overview = overview_response.data
        alice_income = sum(Income.objects.filter(
            user=self.alice, date__year=2026, date__month=9
        ).values_list('amount', flat=True))
        alice_expenses = sum(Expense.objects.filter(
            user=self.alice, date__year=2026, date__month=9
        ).values_list('amount', flat=True))
        self.assertEqual(overview['total_income'], f'{alice_income:.2f}')
        self.assertEqual(overview['total_expenses'], f'{alice_expenses:.2f}')
        category_data = {item['category']: item for item in overview['categories']}
        alice_food_spend = sum(Expense.objects.filter(
            user=self.alice,
            category='Food',
            date__year=2026,
            date__month=9,
        ).values_list('amount', flat=True))
        self.assertEqual(category_data['Food']['total_amount'], f'{alice_food_spend:.2f}')
        self.assertEqual(overview['savings_goals']['completed_goals_count'], 1)

        report_response = self.client.post('/api/reports/generate/', {
            'month': '2026-09',
        }, format='json')
        self.assertEqual(report_response.status_code, status.HTTP_201_CREATED)
        report_data = report_response.data['data']
        self.assertEqual(report_data['summary']['total_income'], f'{alice_income:.2f}')
        self.assertEqual(report_data['summary']['total_expenses'], f'{alice_expenses:.2f}')
        self.assertEqual(report_data['summary']['total_budgeted'], '6500.00')
        report_food = next(item for item in report_data['budget_breakdown'] if item['category'] == 'Food')
        self.assertEqual(report_food['actual_spent'], f'{alice_food_spend:.2f}')
        report_goal = next(item for item in report_data['savings_goals'] if item['id'] == goal_id)
        self.assertTrue(report_goal['is_completed'])
        self.assertEqual(report_goal['progress_percentage'], 100.0)

        self.auth_bob()
        bob_overview_response = self.client.get('/api/analytics/overview/?month=2026-09')
        self.assertEqual(bob_overview_response.status_code, status.HTTP_200_OK)
        self.assertEqual(bob_overview_response.data['total_income'], '30000.00')
        self.assertEqual(bob_overview_response.data['total_expenses'], '8000.00')
        bob_notifications_response = self.client.get('/api/notifications/')
        self.assertEqual(bob_notifications_response.status_code, status.HTTP_200_OK)
        bob_notification_ids = {item['id'] for item in bob_notifications_response.data}
        self.assertNotIn(
            Notification.objects.get(user=self.alice, related_context=alert_context).id,
            bob_notification_ids,
        )
        self.assertNotIn(
            Notification.objects.get(user=self.alice, related_context=f'goal:{goal_id}:completed').id,
            bob_notification_ids,
        )

    def test_analytics_date_range_filtering(self):
        """Analytics endpoint respects start_date and end_date filtering"""
        self.auth_alice()
        # Filter range covering only first income and first expense: 2026-09-01 to 2026-09-12
        response = self.client.get('/api/analytics/overview/?start_date=2026-09-01&end_date=2026-09-12')
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        data = response.data
        self.assertEqual(data['total_income'], '5000.00') # only Pocket Money
        self.assertEqual(data['total_expenses'], '3000.00') # only Food
        self.assertEqual(data['net_savings'], '2000.00')
        self.assertEqual(data['savings_rate'], 40.0)

    def test_analytics_empty_and_zero_values_safety(self):
        """User with no records receives clean zeroed structures without 500 or division errors"""
        new_user = User.objects.create_user(username='zero_user', email='zero@example.com', password='Password123!')
        token = str(RefreshToken.for_user(new_user).access_token)
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

        response = self.client.get('/api/analytics/overview/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['total_income'], '0.00')
        self.assertEqual(response.data['total_expenses'], '0.00')
        self.assertEqual(response.data['net_savings'], '0.00')
        self.assertEqual(response.data['savings_rate'], 0.0)

    def test_analytics_monthly_trends(self):
        """Monthly trends endpoint returns chronological monthly summaries"""
        self.auth_alice()
        response = self.client.get('/api/analytics/trends/?months=6')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('trends', response.data)
        self.assertEqual(len(response.data['trends']), 6)

    def test_analytics_monthly_trends_end_on_selected_month(self):
        """Trend charts can be anchored to the selected month, including year boundaries."""
        self.auth_alice()
        response = self.client.get('/api/analytics/trends/?months=6&month=2026-01')

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            [item['month'] for item in response.data['trends']],
            ['2025-08', '2025-09', '2025-10', '2025-11', '2025-12', '2026-01'],
        )

    def test_analytics_user_isolation(self):
        """Alice only retrieves her own analytics; Bob's numbers are isolated"""
        self.auth_alice()
        res_alice = self.client.get('/api/analytics/overview/?month=2026-09')
        self.assertEqual(res_alice.data['total_income'], '20000.00')

        self.auth_bob()
        res_bob = self.client.get('/api/analytics/overview/?month=2026-09')
        self.assertEqual(res_bob.data['total_income'], '30000.00')
        self.assertEqual(res_bob.data['total_expenses'], '8000.00')

    # --- 2. Report Generation Workflow ---
    def test_generate_monthly_financial_report(self):
        """POST /api/reports/generate/ creates report snapshot matching actual records and triggers notification"""
        self.auth_alice()
        initial_notifs = Notification.objects.filter(user=self.alice).count()

        response = self.client.post('/api/reports/generate/', {'month': '2026-09'}, format='json')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)

        report_id = response.data['id']
        report = Report.objects.get(id=report_id)
        self.assertEqual(report.user, self.alice)
        self.assertEqual(report.data['summary']['total_income'], '20000.00')
        self.assertEqual(report.data['summary']['total_expenses'], '5000.00')
        self.assertEqual(report.data['summary']['net_savings'], '15000.00')
        self.assertEqual(report.data['summary']['total_budgeted'], '6500.00')
        self.assertEqual(len(report.data['expense_transactions']), 2)
        self.assertEqual(len(report.data['income_transactions']), 2)
        self.assertEqual(report.data['savings_goals'][0]['name'], 'Emergency Fund')
        self.assertEqual(report.data['savings_goals'][0]['target_amount'], '50000.00')
        self.assertEqual(report.data['savings_goals'][0]['current_amount'], '10000.00')
        self.assertEqual(report.data['savings_goals'][0]['progress_percentage'], 20.0)
        self.assertFalse(report.data['savings_goals'][0]['is_completed'])

        food_budget = next(item for item in report.data['budget_breakdown'] if item['category'] == 'Food')
        self.assertEqual(food_budget['allocated_budget'], '4000.00')
        self.assertEqual(food_budget['actual_spent'], '3000.00')
        self.assertEqual(food_budget['remaining'], '1000.00')

        self.assertEqual(
            sum(Decimal(item['amount']) for item in report.data['income_transactions']),
            Decimal('20000.00')
        )
        self.assertEqual(
            sum(Decimal(item['amount']) for item in report.data['expense_transactions']),
            Decimal('5000.00')
        )

        # Verification of notification created
        self.assertEqual(Notification.objects.filter(user=self.alice).count(), initial_notifs + 1)
        latest_notif = Notification.objects.filter(user=self.alice).order_by('-id').first()
        self.assertIn('Report Ready', latest_notif.title)

    def test_generating_same_month_refreshes_report_instead_of_duplicating(self):
        self.auth_alice()
        first_response = self.client.post(
            '/api/reports/generate/',
            {'month': '2026-09'},
            format='json',
        )
        self.assertEqual(first_response.status_code, status.HTTP_201_CREATED)
        first_report_id = first_response.data['id']
        first_generated_at = Report.objects.get(id=first_report_id).generated_at

        Expense.objects.create(
            user=self.alice,
            category='Food',
            amount=Decimal('250.00'),
            date=date(2026, 9, 28),
        )
        second_response = self.client.post(
            '/api/reports/',
            {'month': '2026-09-30'},
            format='json',
        )

        self.assertEqual(second_response.status_code, status.HTTP_200_OK)
        self.assertEqual(second_response.data['id'], first_report_id)
        self.assertEqual(Report.objects.filter(user=self.alice).count(), 1)
        self.assertEqual(
            Report.objects.get(id=first_report_id).data['summary']['total_expenses'],
            '5250.00',
        )
        self.assertGreaterEqual(
            Report.objects.get(id=first_report_id).generated_at,
            first_generated_at,
        )
        self.assertEqual(
            Notification.objects.filter(
                user=self.alice,
                related_context=f'report:{first_report_id}',
            ).count(),
            1,
        )

    def test_invalid_report_month_is_rejected_without_creating_records(self):
        self.auth_alice()

        for month in ('2026-13', '2026-09-31'):
            with self.subTest(month=month):
                response = self.client.post(
                    '/api/reports/generate/',
                    {'month': month},
                    format='json',
                )
                self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

        self.assertFalse(Report.objects.filter(user=self.alice).exists())
        self.assertFalse(
            Notification.objects.filter(
                user=self.alice,
                notification_type='system',
                title__startswith='Report Ready:',
            ).exists()
        )

    def test_report_export_csv_and_json(self):
        """GET /api/reports/<id>/export/ streams CSV, XLSX, PDF, or JSON data"""
        self.auth_alice()
        gen_res = self.client.post('/api/reports/generate/', {'month': '2026-09'}, format='json')
        report_id = gen_res.data['id']

        # 1. Export as CSV
        res_csv = self.client.get(f'/api/reports/{report_id}/export/?format=csv')
        self.assertEqual(res_csv.status_code, status.HTTP_200_OK)
        self.assertEqual(res_csv['Content-Type'], 'text/csv')
        self.assertIn('attachment;', res_csv['Content-Disposition'])
        csv_text = res_csv.content.decode('utf-8')
        self.assertIn('BUDGETBUDDY MONTHLY FINANCIAL REPORT', csv_text)
        self.assertIn('20000.00', csv_text)
        self.assertIn('5000.00', csv_text)
        self.assertIn('₹20000.00', csv_text)
        self.assertNotIn('$', csv_text)

        # 2. Export as JSON
        res_json = self.client.get(f'/api/reports/{report_id}/export/?format=json')
        self.assertEqual(res_json.status_code, status.HTTP_200_OK)
        self.assertEqual(res_json.data['summary']['total_income'], '20000.00')

        # 3. Export as XLSX and verify it is a readable workbook
        res_xlsx = self.client.get(f'/api/reports/{report_id}/export/?format=xlsx')
        self.assertEqual(res_xlsx.status_code, status.HTTP_200_OK)
        self.assertEqual(
            res_xlsx['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('.xlsx', res_xlsx['Content-Disposition'])
        workbook = load_workbook(BytesIO(res_xlsx.content), read_only=True)
        self.assertIn('Summary', workbook.sheetnames)
        self.assertIn('Budgets', workbook.sheetnames)
        self.assertIn('Savings Goals', workbook.sheetnames)
        self.assertEqual(workbook['Summary']['B5'].value, 'Value')
        workbook.close()

        # 4. Export as PDF and verify the document signature and filename
        res_pdf = self.client.get(f'/api/reports/{report_id}/export/?format=pdf')
        self.assertEqual(res_pdf.status_code, status.HTTP_200_OK)
        self.assertEqual(res_pdf['Content-Type'], 'application/pdf')
        self.assertIn('.pdf', res_pdf['Content-Disposition'])
        self.assertTrue(res_pdf.content.startswith(b'%PDF-'))


    def test_report_user_isolation(self):
        """User cannot view, export, or delete another user's report"""
        self.auth_alice()
        gen_res = self.client.post('/api/reports/generate/', {'month': '2026-09'}, format='json')
        alice_report_id = gen_res.data['id']

        # Bob attempts to view, export, delete Alice's report
        self.auth_bob()
        self.assertEqual(self.client.get(f'/api/reports/{alice_report_id}/').status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(self.client.get(f'/api/reports/{alice_report_id}/export/').status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(self.client.delete(f'/api/reports/{alice_report_id}/').status_code, status.HTTP_404_NOT_FOUND)

        # Alice deletes her own report
        self.auth_alice()
        del_res = self.client.delete(f'/api/reports/{alice_report_id}/')
        self.assertEqual(del_res.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(Report.objects.filter(id=alice_report_id).exists())

    def test_unauthenticated_requests_return_401(self):
        """Unauthenticated requests to analytics and report endpoints return 401"""
        self.clear_auth()
        self.assertEqual(self.client.get('/api/analytics/overview/').status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(self.client.get('/api/analytics/trends/').status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(self.client.get('/api/reports/').status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(self.client.post('/api/reports/generate/', {}).status_code, status.HTTP_401_UNAUTHORIZED)
